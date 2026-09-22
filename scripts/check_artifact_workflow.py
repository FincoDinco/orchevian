"""Live local-model file-creation acceptance using an isolated fixture library.

Run with --model backend/name --output NEW_DIRECTORY. Files and reports stay in
that directory; the user's library is never opened and no models are downloaded.
Native files and previews require human layout/content review after this check.
"""

import argparse
import hashlib
import io
import json
import multiprocessing
import re
import threading
import zipfile
from datetime import UTC, datetime
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image

from llm_engine.artifacts.generators import validate_output
from llm_engine.artifacts.specs import Spec
from llm_engine.backends.registry import BackendRegistry
from llm_engine.domain.models import GenerationParams
from llm_engine.services.artifacts import ArtifactService
from llm_engine.services.chat import ChatService
from llm_engine.services.documents import DocumentService
from llm_engine.services.session import ModelSession
from llm_engine.store.library import LibraryService
from llm_engine.store.sqlite import SqliteStore

BRIEF = """Alder Community Workshop
Event date: October 24, 2026. Expected participants: 40.
Purpose: teach residents basic bicycle maintenance.
Budget in USD: Room rental 600; Materials 400; Refreshments 200. Total: 1200.
Program: Prepare the room, Run the workshop, Collect feedback.
Coordinator: Morgan Lee. Registration collects full name and consent to attend.
"""
PRIMARY = """Use the project brief to create exactly four files:
1. proposal.docx: a concise proposal with event date, participants, coordinator,
purpose, program steps, and a table of the three budget items and total.
2. summary.pdf: a one-page report with the same key facts and budget total.
3. budget.xlsx: one Budget sheet, columns Item and Cost (USD), three expense rows
in the brief's order, then a Total row using SUM. Format the costs as USD and
include a bar chart. Keep costs as numbers.
4. briefing.pptx: exactly three editable slides on overview, budget, and program;
include speaker notes and identify the source brief in the notes.
Use the supplied facts and cite the brief. Do not invent extra facts.
"""
EXTRA = """Use the project brief to create exactly four files:
registration.pdf: a fillable form with a text field named full_name and a checkbox
named consent. Use these exact internal field names;
report.html: a standalone report with the event details and complete budget;
budget.json: structured budget data with numeric costs and total;
process.svg: an ordered diagram of the three program steps, in their given order.
"""


class VisibleHTML(HTMLParser):
    """Inspect rendered body text and real table cells, excluding head metadata."""

    def __init__(self, text):
        super().__init__(convert_charrefs=True)
        self.in_body = False
        self.hidden = 0
        self.parts = []
        self.rows = []
        self.row = None
        self.cell = None
        self.feed(text)

    def handle_starttag(self, tag, attrs):
        if tag == "body":
            self.in_body = True
        if tag in {"script", "style"}:
            self.hidden += 1
        if self.in_body and not self.hidden:
            if tag == "tr":
                self.row = []
            elif tag in {"td", "th"}:
                self.cell = []

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self.cell is not None:
            if self.row is not None:
                self.row.append("".join(self.cell).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        if tag == "body":
            self.in_body = False

    def handle_data(self, data):
        if self.in_body and not self.hidden:
            self.parts.append(data)
            if self.cell is not None:
                self.cell.append(data)


PROGRAM = ["Prepare the room", "Run the workshop", "Collect feedback"]
BUDGET = [("Room rental", 600), ("Materials", 400), ("Refreshments", 200), ("Total", 1200)]


def contains_fact(text, fact):
    return bool(re.search(rf"(?<!\w){re.escape(fact)}(?!\w)", text, re.I))


def budget_problems(rows, label):
    problems = []
    for name, amount in BUDGET:
        row = next((r for r in rows if r and str(r[0]).strip().casefold() == name.casefold()), [])
        value = str(row[1]).strip().replace("$", "").replace(",", "") if len(row) > 1 else ""
        if value not in {str(amount), f"{amount:.2f}"}:
            problems.append(f"{label}: {name} table amount is not {amount}")
    return problems


def json_budget_problems(data):
    """Allow nested model-chosen keys, but require labels paired with numeric costs."""
    found = set()

    def visit(value):
        if isinstance(value, dict):
            labels = [str(k).casefold().replace("_", " ") for k in value]
            labels += [v.casefold() for v in value.values() if isinstance(v, str)]
            numbers = [v for v in value.values() if type(v) in {int, float}]
            for name, amount in BUDGET:
                key = name.casefold()
                # Direct mappings associate each key with its own value.
                if any(str(k).casefold().replace("_", " ") == key
                       and type(v) in {int, float} and v == amount for k, v in value.items()):
                    found.add(name)
                # A labelled item object can use cost, amount, value, etc.
                elif key in labels and amount in numbers and len(numbers) == 1:
                    found.add(name)
            for child in value.values():
                visit(child)
        elif isinstance(value, list):
            for child in value:
                visit(child)

    visit(data)
    return [f"JSON: missing numeric {name} amount {amount}"
            for name, amount in BUDGET if name not in found]


def form_problems(path):
    """Verify canonical fields and page widgets survive filling and reopening."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import NameObject

    problems = []
    reader = PdfReader(path)
    fields = reader.get_fields() or {}
    for name, kind in (("full_name", "/Tx"), ("consent", "/Btn")):
        if fields.get(name, {}).get("/FT") != kind:
            problems.append(f"PDF {name} field is missing or has the wrong type")
    if problems:
        return problems
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    values = {"full_name": "Acceptance Test", "consent": NameObject("/Yes")}
    writer.update_page_form_field_values(None, values, auto_regenerate=False)
    stream = io.BytesIO()
    writer.write(stream)
    filled = PdfReader(io.BytesIO(stream.getvalue()))
    fields = filled.get_fields() or {}
    widgets = {}
    for page in filled.pages:
        for ref in page.get("/Annots", []):
            widget = ref.get_object()
            if widget.get("/Subtype") == "/Widget":
                field = widget.get("/Parent", widget).get_object()
                widgets[field.get("/T")] = (widget, field)
    for name, expected in values.items():
        if fields.get(name, {}).get("/V") != expected or name not in widgets:
            problems.append(f"PDF {name} value or page widget did not survive filling")
            continue
        widget, field = widgets[name]
        if widget.get("/V", field.get("/V")) != expected or not widget.get("/AP", {}).get("/N"):
            problems.append(f"PDF {name} widget value or appearance is missing")
        if name == "consent" and widget.get("/AS") != expected:
            problems.append("PDF checkbox appearance does not match its filled value")
    return problems


def check_fixture_content(output, *, require_revision=True):
    """Independent read-back of this fixture's facts and requested native features."""
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    from pypdf import PdfReader

    problems = []
    stages = [("primary", 40), ("revision", 50)] if require_revision else [("primary", 40)]
    for stage, participants in stages:
        doc = Document(output / stage / "proposal.docx")
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        rows = [tuple(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        if not re.search(rf"\b{participants}\b", paragraphs):
            problems.append(f"{stage}: participant count missing")
        for fact in ("October 24, 2026", "Morgan Lee", "bicycle maintenance"):
            if fact not in paragraphs:
                problems.append(f"{stage}: missing {fact}")
        problems.extend(budget_problems(rows, stage))
        if any(step.casefold() not in paragraphs.casefold() for step in PROGRAM):
            problems.append(f"{stage}: program steps missing")
        if not re.search(r"brief|P[a-f0-9]+v\d+", paragraphs, re.I):
            problems.append(f"{stage}: source brief citation missing")
    pdf = PdfReader(output / "primary" / "summary.pdf")
    pdf_text = "\n".join(p.extract_text() for p in pdf.pages)
    if len(pdf.pages) != 1 or "=SUM(" in pdf_text.upper():
        problems.append("PDF must be one page with numeric totals, not formula text")
    for fact in ("October 24, 2026", "40", "Morgan Lee", "1200"):
        if not contains_fact(pdf_text.replace(",", ""), fact.replace(",", "")):
            problems.append(f"PDF missing {fact}")
    for name, amount in BUDGET:
        if not re.search(rf"{name}\s+\$?{amount}(?:\.00)?\b",
                         pdf_text.replace(",", ""), re.I):
            problems.append(f"PDF budget missing {name}: {amount}")
    if any(not contains_fact(pdf_text, step) for step in PROGRAM):
        problems.append("PDF program steps missing")
    book = load_workbook(output / "primary" / "budget.xlsx")
    try:
        sheet = book["Budget"]
        if [sheet[f"B{n}"].value for n in range(2, 5)] != [600, 400, 200]:
            problems.append("Spreadsheet costs do not match numeric source values")
        if str(sheet["B5"].value).upper().replace("$", "") != "=SUM(B2:B4)":
            problems.append("Spreadsheet total does not sum the three expense rows")
        if not sheet._charts or sheet._charts[0].__class__.__name__ != "BarChart":
            problems.append("Spreadsheet bar chart is missing")
        if book.sheetnames != ["Budget"]:
            problems.append("Spreadsheet must contain only the Budget sheet")
        if [sheet[f"A{n}"].value for n in range(2, 6)] != [name for name, _ in BUDGET]:
            problems.append("Spreadsheet budget labels are missing or out of order")
        if any("$" not in sheet[f"B{n}"].number_format for n in range(2, 6)):
            problems.append("Spreadsheet currency formatting is missing")
    finally:
        book.close()
    deck = Presentation(output / "primary" / "briefing.pptx")
    if len(deck.slides) != 3 or any(not slide.notes_slide.notes_text_frame.text.strip()
                                  for slide in deck.slides):
        problems.append("Deck needs three slides with speaker notes")
    slide_text = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                  for slide in deck.slides]
    for fact in ("October 24, 2026", "40", "Morgan Lee", "bicycle maintenance", *PROGRAM):
        if not contains_fact("\n".join(slide_text), fact):
            problems.append(f"Deck missing {fact}")
    if len(slide_text) == 3:
        for name, amount in BUDGET:
            if not re.search(rf"{name}[^\n]*\b{amount}\b",
                             slide_text[1].replace(",", ""), re.I):
                problems.append(f"Deck budget missing {name}: {amount}")
        if any(step.casefold() not in slide_text[2].casefold() for step in PROGRAM):
            problems.append("Deck program slide missing steps")
    if any(not re.search(r"brief|P[a-f0-9]+v\d+", slide.notes_slide.notes_text_frame.text, re.I)
           for slide in deck.slides):
        problems.append("Deck notes must identify the source brief")
    if (output / "extended").exists():
        problems.extend(form_problems(output / "extended" / "registration.pdf"))
        from defusedxml import ElementTree

        diagram = ElementTree.parse(output / "extended" / "process.svg")
        labels = ["".join(node.itertext()) for node in diagram.iter("{http://www.w3.org/2000/svg}text")]
        if [text for text in labels if text in PROGRAM] != PROGRAM:
            problems.append("Diagram did not preserve the program order")
        html = VisibleHTML((output / "extended" / "report.html").read_text(encoding="utf-8"))
        visible = "\n".join(html.parts)
        if re.search(r"</?(?:h[1-6]|p|table|tr|td|th|strong|em|div)\b", visible, re.I):
            problems.append("HTML displays literal markup instead of formatted content")
        for fact in ("October 24, 2026", "40", "Morgan Lee", "bicycle maintenance"):
            if not contains_fact(visible, fact):
                problems.append(f"HTML missing {fact}")
        problems.extend(budget_problems(html.rows, "HTML"))
        problems.extend(json_budget_problems(json.loads(
            (output / "extended" / "budget.json").read_text(encoding="utf-8"),
        )))
    return problems


def run(model_id, output, *, timeout=300, max_tokens=4096, extended=False):
    output.mkdir(parents=True, exist_ok=False)
    report = {
        "ok": False, "model": model_id, "turns": [], "files": [], "checks": [],
        "started_at": datetime.now(UTC).isoformat(),
        "settings": {"timeout": timeout, "max_tokens": max_tokens, "extended": extended},
        "review_limits": "Automated content/structure checks; native Office layout and "
                         "spreadsheet recalculation require separate review.",
    }
    registry = BackendRegistry()
    session = ModelSession(registry)
    original_generate = session.generate
    attempt = 0

    def traced_generate(messages, params, cancel):
        nonlocal attempt
        attempt += 1
        prefix = output / f"model-attempt-{attempt}"
        prefix.with_suffix(".prompt.json").write_text(json.dumps([
            {"role": turn.role, "content": turn.content} for turn in messages
        ], indent=2), encoding="utf-8")
        parts = []
        stream = original_generate(messages, params, cancel)
        try:
            for token in stream:
                parts.append(token)
                yield token
        finally:
            stream.close()
            prefix.with_suffix(".response.txt").write_text("".join(parts), encoding="utf-8")

    session.generate = traced_generate
    snapshots = {}
    try:
        models, availability = registry.list_models()
        model = next((m for m in models if m.ref.id == model_id), None)
        if model is None:
            raise ValueError(f"Model not installed: {model_id}")
        available, reason = availability[str(model.ref.backend)]
        if not available:
            raise ValueError(reason or "Runtime unavailable")
        with SqliteStore(output / "fixture.db") as store:
            library = LibraryService(store)
            errors, progress = [], []
            chat = ChatService(library, session,
                               on_error=lambda cid, error: errors.append(str(error)))
            chat.on_artifact_progress = lambda cid, message: (
                progress.append(message), print(message, flush=True),
            )
            project = library.create_project("Artifact acceptance fixture")
            brief = output / "brief.txt"
            brief.write_text(BRIEF, encoding="utf-8")
            chat.documents.import_project_file(project.id, brief, threading.Event())
            cid = library.create_conversation(project_id=project.id, model=model.ref).summary.id

            def turn(name, conversation_id, prompt, request):
                errors.clear()
                progress.clear()
                print(f"Starting {name}", flush=True)
                before = {a.id for a in chat.artifacts.list(conversation_id)}
                chat.send(conversation_id, prompt,
                          GenerationParams(temperature=0, max_tokens=max_tokens),
                          artifact_request=request)
                chat._worker_thread.join(timeout)
                if chat._worker_thread.is_alive():
                    chat.cancel_current()
                    session.force_unload()
                    chat._worker_thread.join(10)
                    errors.append(f"Turn exceeded {timeout} seconds")
                record = {"name": name, "errors": list(errors), "progress": list(progress)}
                report["turns"].append(record)
                if errors:
                    raise RuntimeError("; ".join(errors))
                files = [a for a in chat.artifacts.list(conversation_id) if a.id not in before]
                folder = output / name
                folder.mkdir()
                for artifact in files:
                    generated, sources = chat.artifacts.get(conversation_id, artifact.id)
                    validate_output(Spec.model_validate(generated.spec), generated.data)
                    chat.artifacts.export(conversation_id, [artifact.id], folder / artifact.name)
                    (folder / (artifact.name + ".spec.json")).write_text(
                        json.dumps(generated.spec, indent=2), encoding="utf-8",
                    )
                    for index, preview in enumerate(generated.previews, 1):
                        with Image.open(io.BytesIO(preview)) as image:
                            image.verify()
                        (folder / f"{artifact.name}.page-{index}.png").write_bytes(preview)
                    snapshots[(conversation_id, artifact.id)] = hashlib.sha256(
                        generated.data,
                    ).hexdigest()
                    report["files"].append({
                        "name": artifact.name, "version": artifact.version, "stage": name,
                        "bytes": len(generated.data), "preview_pages": len(generated.previews),
                        "source_snapshots": sources, "warning": generated.warning,
                    })
                record["files"] = [a.name for a in files]
                return files

            try:
                primary = turn("primary", cid, PRIMARY, {})
                required = {"proposal.docx", "summary.pdf", "budget.xlsx", "briefing.pptx"}
                if {a.name for a in primary} != required:
                    raise ValueError("Primary batch did not produce exactly the requested files")
                if not all(chat.artifacts.get(cid, a.id)[1] for a in primary):
                    raise ValueError("Primary files lost their project-source snapshots")
                report["checks"].append("primary files, native read-back, previews and sources")
                chat.artifacts.export(cid, [a.id for a in primary], output / "primary-batch.zip")
                with zipfile.ZipFile(output / "primary-batch.zip") as archive:
                    if set(archive.namelist()) != required:
                        raise ValueError("Batch export omitted a requested file")
                report["checks"].append("multi-file ZIP export")
                proposal = next(a for a in primary if a.name == "proposal.docx")
                revised = turn("revision", cid,
                               "Revise only proposal.docx: change expected participants from 40 "
                               "to 50. Keep the event date, all three costs, total, coordinator "
                               "and program unchanged. Keep the filename proposal.docx.",
                               {"revision": proposal.id})
                if len(revised) != 1 or revised[0].name != proposal.name or revised[0].version != 2:
                    raise ValueError("Revision did not retain its filename and increment version")
                chat.artifacts.add_to_project(cid, revised[0].id, project.id)
                check_cid = library.create_conversation(project_id=project.id).summary.id
                store.add_message(check_cid, "user", "Workshop participants and budget")
                context = DocumentService(store).context(check_cid, "workshop participants budget")
                (output / "project-reuse-context.txt").write_text(context, encoding="utf-8")
                if "proposal.docx" not in context:
                    raise ValueError("Revised project copy was not available for retrieval")
                report["checks"].append("revision versioning and project-copy retrieval")
                if extended:
                    # Keep the original brief separate from the revised project copy:
                    # two conflicting participant counts would make this fixture ambiguous.
                    extended_project = library.create_project("Extended artifact fixture")
                    chat.documents.import_project_file(
                        extended_project.id, brief, threading.Event(),
                    )
                    other = library.create_conversation(project_id=extended_project.id,
                                                        model=model.ref).summary.id
                    files = turn("extended", other, EXTRA, {})
                    if {a.name for a in files} != {
                        "registration.pdf", "report.html", "budget.json", "process.svg",
                    }:
                        raise ValueError("Extended batch omitted requested files")
                    report["checks"].append("extended files, native structure and preview decoding")
            finally:
                chat.cancel_current()
                session.force_unload()
                if chat._worker_thread:
                    chat._worker_thread.join(10)
        with SqliteStore(output / "fixture.db") as reopened:
            artifacts = ArtifactService(reopened)
            for (cid, key), digest in snapshots.items():
                if hashlib.sha256(artifacts.get(cid, key)[0].data).hexdigest() != digest:
                    raise ValueError("A retained file changed after reopening the library")
        report["checks"].append("original and revised bytes survive library reopen")
        report["content_problems"] = check_fixture_content(output)
        report["ok"] = not report["content_problems"]
        if report["ok"]:
            report["checks"].append("fixture facts, numeric totals, formatting and native features")
    except Exception as exc:
        report["error"] = str(exc)
    finally:
        session.force_unload()
        registry.close()
        (output / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2), flush=True)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", required=True)
    parser.add_argument("--output", type=Path, required=True,
                        help="New directory for fixtures/output")
    parser.add_argument("--timeout", type=float, default=300)
    parser.add_argument("--max-tokens", type=int, default=4096)
    parser.add_argument("--extended", action="store_true")
    args = parser.parse_args()
    if args.timeout <= 0 or args.max_tokens <= 0:
        parser.error("timeout and max-tokens must be positive")
    report = run(args.model, args.output, timeout=args.timeout, max_tokens=args.max_tokens,
                 extended=args.extended)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    multiprocessing.freeze_support()
    raise SystemExit(main())
