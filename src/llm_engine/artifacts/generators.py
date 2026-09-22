"""Native format writers and read-back validation, executed in a disposable process."""

from __future__ import annotations

import csv
import html
import io
import json
import math
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from importlib.util import find_spec

from llm_engine.artifacts.specs import Spec


@dataclass(frozen=True)
class Capability:
    format: str
    kinds: tuple[str, ...]
    dependency: str | None
    features: str
    preview: str = "Content preview; open the file to inspect native layout"

    @property
    def available(self):
        return not self.missing_dependencies

    @property
    def missing_dependencies(self):
        required = {"reportlab", "pypdfium2", "PIL"}
        if self.dependency:
            required.add(self.dependency)
        return sorted(name for name in required if find_spec(name) is None)


REGISTRY = {
    item.format: item
    for item in (
        Capability("docx", ("document", "form"), "docx", "Headings, lists, tables, page breaks"),
        Capability(
            "pdf",
            ("document", "slides", "form", "chart", "diagram"),
            "reportlab",
            "Paginated documents, slide handouts, fillable text/checkbox fields",
            "PDF pages",
        ),
        Capability("rtf", ("document",), None, "Unicode plain text"),
        Capability("txt", ("document",), None, "Plain text"),
        Capability("md", ("document",), None, "Headings, lists and tables"),
        Capability(
            "xlsx",
            ("spreadsheet",),
            "openpyxl",
            "Typed cells, sheets, number formats, formulas, bar/line charts",
        ),
        Capability("csv", ("spreadsheet",), None, "One sheet of values"),
        Capability("tsv", ("spreadsheet",), None, "One sheet of values"),
        Capability("pptx", ("slides",), "pptx", "Editable titles/bullets and speaker notes"),
        Capability("html", ("document",), None, "Standalone styled report, no external assets"),
        Capability("json", ("data",), None, "JSON objects and arrays"),
        Capability("yaml", ("data",), "yaml", "YAML objects and arrays"),
        Capability("xml", ("data",), None, "Typed data tree under a data element"),
        Capability(
            "svg",
            ("chart", "diagram"),
            "reportlab",
            "Bar/line charts, process steps",
            "Vector drawing",
        ),
        Capability(
            "png", ("chart", "diagram"), "reportlab", "Bar/line charts, process steps", "Image"
        ),
    )
}


@dataclass(frozen=True)
class Generated:
    spec: dict
    data: bytes
    previews: tuple[bytes, ...]
    text: str
    warning: str


def check_spec(spec: Spec):
    cap = REGISTRY[spec.format]
    if not cap.available:
        raise ValueError(f"{spec.format.upper()} requires {', '.join(cap.missing_dependencies)}")
    if spec.kind not in cap.kinds:
        raise ValueError(f"{spec.format.upper()} supports: {', '.join(cap.kinds)}")
    if spec.format in {"csv", "tsv"} and len(spec.sheets) != 1:
        raise ValueError("CSV/TSV can contain exactly one sheet; choose XLSX for multiple sheets")

    # All input text is bounded even for nested structured-data values.
    def visit(value, depth=0):
        if depth > 20:
            raise ValueError("Structured data exceeds the nesting limit")
        if isinstance(value, str) and (
            len(value) > 8000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)
        ):
            raise ValueError("Text contains control characters or exceeds 8,000 characters")
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("Numbers must be finite")
        if isinstance(value, dict):
            for key, item in value.items():
                visit(key, depth + 1)
                visit(item, depth + 1)
        elif isinstance(value, list):
            for item in value:
                visit(item, depth + 1)

    visit(spec.model_dump())
    if len(spec.model_dump_json()) > 120_000:
        raise ValueError("File specification exceeds 120,000 characters")


def _document_blocks(spec):
    # Small models often repeat the title as their first heading. The writers
    # already render spec.title; suppress only that exact leading duplicate.
    for index, block in enumerate(spec.blocks):
        if (index == 0 and spec.title.strip() and block.type == "heading"
                and re.sub(r"^#{1,6}\s+", "", block.text.strip()).casefold()
                == spec.title.strip().casefold()):
            continue
        yield block


def text_content(spec):
    lines = [spec.title]
    for block in _document_blocks(spec):
        lines.append(block.text)
        lines.extend(block.items)
        lines.extend(" | ".join(str(c) if c is not None else "" for c in row) for row in block.rows)
    for sheet in spec.sheets:
        lines.append(f"Sheet: {sheet.name}")
        lines.extend(" | ".join(str(c) if c is not None else "" for c in row) for row in sheet.rows)
    for slide in spec.slides:
        lines.extend([slide.title, *slide.bullets, "Speaker notes: " + slide.notes])
    lines.extend(
        f"{field.label}: {'[ ]' if field.type == 'checkbox' else '________'}"
        for field in spec.fields
    )
    if spec.data is not None:
        lines.append(json.dumps(spec.data, ensure_ascii=False, indent=2))
    if spec.chart:
        lines.extend(
            f"{label}: {value}" for label, value in zip(spec.chart.labels, spec.chart.values)
        )
    lines.extend(spec.steps)
    return "\n".join(line for line in lines if line)


def _styles():
    from pathlib import Path

    import reportlab
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("Artifact", "Vera.ttf"), ("ArtifactBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = (
            "ArtifactBold" if style.name.startswith(("Title", "Heading")) else "Artifact"
        )
        style.wordWrap = "CJK"
    styles["BodyText"].spaceAfter = 8
    return styles


def _drawing(spec):
    from textwrap import wrap

    from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
    from reportlab.lib import colors

    _styles()
    drawing = Drawing(720, 440 if spec.kind == "chart" else 100 + 68 * len(spec.steps))
    for index, line in enumerate(wrap(spec.title, 70)):
        drawing.add(
            String(24, drawing.height - 32 - index * 19, line, fontName="ArtifactBold", fontSize=16)
        )
    if spec.kind == "diagram":
        for i, step in enumerate(spec.steps):
            y = drawing.height - 100 - i * 68
            drawing.add(
                Rect(
                    70,
                    y,
                    580,
                    44,
                    fillColor=colors.HexColor("#edf3fc"),
                    strokeColor=colors.HexColor("#3866ad"),
                    rx=6,
                    ry=6,
                )
            )
            # Fit two lines without accepting arbitrary drawing coordinates.
            for j, line in enumerate(wrap(step, 65)):
                drawing.add(String(90, y + 27 - j * 14, line, fontName="Artifact", fontSize=11))
            if i < len(spec.steps) - 1:
                drawing.add(Line(360, y, 360, y - 20))
                drawing.add(
                    Polygon([355, y - 15, 360, y - 21, 365, y - 15], fillColor=colors.black)
                )
    else:
        from reportlab.graphics.charts.barcharts import VerticalBarChart
        from reportlab.graphics.charts.linecharts import HorizontalLineChart

        chart = VerticalBarChart() if spec.chart.type == "bar" else HorizontalLineChart()
        chart.x, chart.y, chart.width, chart.height = 60, 100, 610, 260
        chart.data = [spec.chart.values]
        chart.categoryAxis.categoryNames = spec.chart.labels
        chart.categoryAxis.labels.fontName = "Artifact"
        chart.categoryAxis.labels.fontSize = 8
        chart.categoryAxis.labels.angle = 35
        chart.categoryAxis.labels.boxAnchor = "ne"
        chart.valueAxis.labels.fontName = "Artifact"
        drawing.add(chart)
    return drawing


def _pdf(spec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas
    from reportlab.platypus import (
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _styles()
    output = io.BytesIO()
    if spec.kind in {"chart", "diagram"}:
        from reportlab.graphics import renderPDF

        renderPDF.drawToFile(_drawing(spec), output)
        return output.getvalue()
    if spec.kind == "form":
        pdf = canvas.Canvas(output, pagesize=A4)
        pdf.setTitle(spec.title)
        for index, field in enumerate(spec.fields):
            if index % 8 == 0:
                if index:
                    pdf.showPage()
                pdf.setFont("ArtifactBold", 18)
                from reportlab.lib.utils import simpleSplit

                for j, line in enumerate(simpleSplit(spec.title, "ArtifactBold", 18, 490)):
                    pdf.drawString(48, 788 - j * 22, line)
            y = 718 - (index % 8) * 82
            pdf.setFont("Artifact", 10)
            from reportlab.lib.utils import simpleSplit

            for j, line in enumerate(simpleSplit(field.label, "Artifact", 10, 490)):
                pdf.drawString(48, y + 16 - j * 12, line)
            args = dict(name=field.name, tooltip=field.label, x=48, y=y - 32)
            if field.type == "checkbox":
                pdf.acroForm.checkbox(**args, size=18, fieldFlags="")
            else:
                pdf.acroForm.textfield(**args, width=490, height=25, maxlen=1000)
        pdf.showPage()
        pdf.save()
        return output.getvalue()
    slide_mode = spec.kind == "slides"
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4) if slide_mode else A4,
        leftMargin=48,
        rightMargin=48,
        topMargin=42,
        bottomMargin=42,
        title=spec.title,
    )
    flow = []

    def para(value, style="BodyText"):
        return Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles[style])

    def table(rows):
        if not rows:
            return
        width = max(len(row) for row in rows)
        if not width:
            return
        grid = [
            [para(c if c is not None else "") for c in row] + [""] * (width - len(row))
            for row in rows
        ]
        item = Table(
            grid, colWidths=[doc.width / width] * width, repeatRows=1, splitByRow=1, splitInRow=1
        )
        item.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf3fc")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#dddddd")),
                ]
            )
        )
        flow.extend([item, Spacer(1, 12)])

    if not slide_mode and spec.title:
        flow.append(para(spec.title, "Title"))
    for block in _document_blocks(spec):
        if block.type == "page_break":
            flow.append(PageBreak())
        elif block.type == "list":
            flow.extend(para("• " + item) for item in block.items)
        elif block.type == "table":
            table(block.rows)
        else:
            flow.append(
                para(block.text, f"Heading{block.level}" if block.type == "heading" else "BodyText")
            )
    for sheet in spec.sheets:
        flow.append(para(sheet.name, "Heading1"))
        # Preview is bounded and explicitly disclosed; native file retains all rows.
        table([row[:8] for row in sheet.rows[:40]])
        if len(sheet.rows) > 40 or any(len(row) > 8 for row in sheet.rows):
            flow.append(para("Preview limited to the first 40 rows and 8 columns."))
    for i, slide in enumerate(spec.slides):
        if i:
            flow.append(PageBreak())
        flow.append(KeepTogether([para(slide.title, "Title"), Spacer(1, 20)]))
        flow.extend(para("• " + bullet, "Heading2") for bullet in slide.bullets)
    if spec.data is not None:
        for line in json.dumps(spec.data, indent=2, ensure_ascii=False).splitlines()[:300]:
            flow.append(para(line))
    doc.build(flow or [para(" ")])
    return output.getvalue()


def _docx(spec):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.font.color.rgb = RGBColor.from_string("243247")
    doc.add_heading(spec.title, 0)
    for block in _document_blocks(spec):
        if block.type == "heading":
            doc.add_heading(block.text, block.level)
        elif block.type == "list":
            for item in block.items:
                doc.add_paragraph(item, "List Bullet")
        elif block.type == "table" and block.rows:
            width = max(map(len, block.rows))
            if width:
                table = doc.add_table(rows=0, cols=width)
                table.style = "Light Shading Accent 1"
                for row in block.rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = "" if value is None else str(value)
        elif block.type == "page_break":
            doc.add_page_break()
        else:
            doc.add_paragraph(block.text)
    for field in spec.fields:
        doc.add_paragraph(field.label, "Heading 2")
        doc.add_paragraph("☐" if field.type == "checkbox" else "_" * 55)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _check_formula(value, sheets, *, cell=None):
    from openpyxl.formula.tokenizer import Token, Tokenizer
    from openpyxl.utils.cell import range_boundaries

    if len(value) > 500 or any(char in value for char in ("[", "]", "|", "\\", "@", ";")):
        raise ValueError("Formula uses unsupported external references or syntax")
    allowed = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF", "ROUND", "ABS"}
    depth = 0
    tokens = Tokenizer(value).items
    if not tokens:
        raise ValueError("Empty spreadsheet formula")
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN":
            if token.value[:-1].upper() not in allowed:
                raise ValueError("Supported formula functions: " + ", ".join(sorted(allowed)))
        if token.subtype == "OPEN":
            depth += 1
        elif token.subtype == "CLOSE":
            depth -= 1
            if depth < 0:
                raise ValueError("Unbalanced formula parentheses")
        if token.type == "OPERAND" and token.subtype == "RANGE":
            reference = token.value
            target_sheet = cell[0] if cell else None
            if "!" in reference:
                sheet, reference = reference.rsplit("!", 1)
                target_sheet = sheet.strip("'").replace("''", "'")
                if target_sheet not in sheets:
                    raise ValueError("Formula refers to an unknown sheet")
            if not re.fullmatch(
                r"\$?[A-Z]{1,2}\$?[1-9][0-9]{0,3}"
                r"(?::\$?[A-Z]{1,2}\$?[1-9][0-9]{0,3})?",
                reference,
            ):
                raise ValueError("Formulas require bounded cell references")
            if cell and target_sheet == cell[0]:
                left, top, right, bottom = range_boundaries(reference)
                if left <= cell[2] <= right and top <= cell[1] <= bottom:
                    from openpyxl.utils import get_column_letter

                    address = f"{cell[0]}!{get_column_letter(cell[2])}{cell[1]}"
                    raise ValueError(
                        f"Formula at {address} references itself. Use the input cells, "
                        "not the formula's output cell, including inside SUM ranges."
                    )
    if depth:
        raise ValueError("Unbalanced formula parentheses")

    # Tokenization alone accepts incomplete expressions such as '=1+'. Validate grammar too.
    tokens = [token for token in tokens if token.type != "WHITE-SPACE"]
    position = 0

    def expression():
        nonlocal position
        atom()
        while position < len(tokens) and tokens[position].type in {Token.OP_IN, Token.OP_POST}:
            operator = tokens[position]
            position += 1
            if operator.type == Token.OP_IN:
                atom()

    def atom():
        nonlocal position
        if position >= len(tokens):
            raise ValueError("Incomplete formula expression")
        token = tokens[position]
        position += 1
        if token.type == Token.OP_PRE:
            atom()
        elif token.type == "OPERAND" and token.subtype in {"RANGE", "NUMBER", "TEXT", "LOGICAL"}:
            return
        elif token.subtype == "OPEN" and token.type in {"PAREN", "FUNC"}:
            count = 1
            expression()
            while position < len(tokens) and tokens[position].type == "SEP":
                if token.type != "FUNC" or tokens[position].value != ",":
                    raise ValueError("Unsupported formula separator")
                position += 1
                count += 1
                expression()
            if position >= len(tokens) or tokens[position].subtype != "CLOSE":
                raise ValueError("Unbalanced formula parentheses")
            position += 1
            expected = {"IF(": (2, 3), "ROUND(": (2,), "ABS(": (1,)}
            if token.value.upper() in expected and count not in expected[token.value.upper()]:
                raise ValueError("Wrong number of formula arguments")
        else:
            raise ValueError("Unsupported or incomplete formula expression")

    expression()
    if position != len(tokens):
        raise ValueError("Unexpected token in formula")


def _xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    book = Workbook()
    book.remove(book.active)
    names = {sheet.name for sheet in spec.sheets}
    for sheet in spec.sheets:
        ws = book.create_sheet(sheet.name)
        for row_number, row in enumerate(sheet.rows, 1):
            for column_number, value in enumerate(row, 1):
                if isinstance(value, str) and value.startswith("="):
                    _check_formula(value, names, cell=(sheet.name, row_number, column_number))
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315B91")
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = 22
        for address, number_format in sheet.number_formats.items():
            ws[address].number_format = number_format
        if sheet.chart != "none":
            chart = BarChart() if sheet.chart == "bar" else LineChart()
            chart.title = spec.title or sheet.name
            chart.add_data(
                Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=ws.max_row),
                titles_from_data=True,
            )
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
            ws.add_chart(chart, f"A{ws.max_row + 3}")
    book.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _pptx(spec):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, item in enumerate(spec.slides, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(1.2))
        title.text_frame.word_wrap = True
        title.text_frame.text = item.title
        title.text_frame.paragraphs[0].font.size = Pt(30)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("244C7C")
        body = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.9))
        body.text_frame.word_wrap = True
        for i, bullet in enumerate(item.bullets):
            paragraph = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
            paragraph.text = "• " + bullet
            paragraph.font.size = Pt(20 if sum(map(len, item.bullets)) < 650 else 16)
            paragraph.space_after = Pt(14)
        footer = slide.shapes.add_textbox(Inches(12), Inches(7), Inches(0.7), Inches(0.3))
        footer.text_frame.text = str(index)
        footer.text_frame.paragraphs[0].font.size = Pt(10)
        slide.notes_slide.notes_text_frame.text = item.notes
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def _markdown(spec):
    lines = ["# " + spec.title, ""] if spec.title else []
    for block in _document_blocks(spec):
        if block.type == "heading":
            lines.append("#" * (block.level + 1) + " " + block.text)
        elif block.type == "list":
            lines.extend("- " + item for item in block.items)
        elif block.type == "table":
            for index, row in enumerate(block.rows):
                lines.append(
                    "| "
                    + " | ".join(str(c).replace("|", "\\|") if c is not None else "" for c in row)
                    + " |"
                )
                if index == 0:
                    lines.append("| " + " | ".join("---" for _ in row) + " |")
        elif block.type == "page_break":
            lines.append("---")
        else:
            lines.append(block.text)
        lines.append("")
    return "\n".join(lines)


def _html(spec):
    esc = html.escape
    parts = [f"<h1>{esc(spec.title)}</h1>"]
    for block in _document_blocks(spec):
        if block.type == "table":
            rows = [
                "<tr>"
                + "".join(f"<td>{esc(str(c) if c is not None else '')}</td>" for c in row)
                + "</tr>"
                for row in block.rows
            ]
            parts.append("<table>" + "".join(rows) + "</table>")
        elif block.type == "list":
            parts.append("<ul>" + "".join(f"<li>{esc(t)}</li>" for t in block.items) + "</ul>")
        elif block.type == "page_break":
            parts.append('<hr style="break-after:page">')
        else:
            tag = f"h{block.level + 1}" if block.type == "heading" else "p"
            parts.append(f"<{tag}>{esc(block.text)}</{tag}>")
    return (
        '<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<meta http-equiv="Content-Security-Policy" '
        "content=\"default-src 'none'; style-src 'unsafe-inline'\">"
        f"<title>{esc(spec.title)}</title><style>"
        "body{font:16px/1.6 system-ui,sans-serif;max-width:900px;margin:48px auto;"
        "padding:0 24px;color:#243247}h1,h2,h3{color:#244c7c}p{white-space:pre-wrap}"
        "table{border-collapse:collapse;width:100%;margin:24px 0}td{border:1px solid #ddd;"
        "padding:8px;overflow-wrap:anywhere}tr:first-child{background:#edf3fc}"
        "@media print{body{margin:0;max-width:none}}"
        "</style></head><body>" + "".join(parts) + "</body></html>"
    )


def _xml(data):
    def fill(element, value):
        element.set("type", type(value).__name__)
        if isinstance(value, dict):
            for key, item in value.items():
                fill(ET.SubElement(element, "entry", key=key), item)
        elif isinstance(value, list):
            for item in value:
                fill(ET.SubElement(element, "item"), item)
        else:
            element.text = json.dumps(value, ensure_ascii=False)

    root = ET.Element("data")
    fill(root, data)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _rtf(text):
    value = []
    for char in text:
        if char in "\\{}":
            value.append("\\" + char)
        elif char == "\n":
            value.append("\\par\n")
        elif ord(char) < 128:
            value.append(char)
        else:
            encoded = char.encode("utf-16-le")
            for i in range(0, len(encoded), 2):
                code = int.from_bytes(encoded[i : i + 2], "little", signed=True)
                value.append(f"\\u{code}?")
    return (
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs22 " + "".join(value) + "}"
    ).encode("ascii")


def pdf_previews(data):
    import pypdfium2 as pdfium

    images = []
    doc = pdfium.PdfDocument(data)
    try:
        doc.init_forms()
        for index in range(min(len(doc), 8)):
            page = doc[index]
            bitmap = page.render(scale=min(1.3, 1200 / max(page.get_size())))
            try:
                output = io.BytesIO()
                bitmap.to_pil().save(output, "PNG")
                images.append(output.getvalue())
            finally:
                bitmap.close()
                page.close()
    finally:
        doc.close()
    return tuple(images)


def validate_output(spec, data):
    """Read native structures back; the extension alone never establishes validity."""
    stream = io.BytesIO(data)
    if spec.format == "docx":
        from docx import Document

        doc = Document(stream)
        if not doc.paragraphs:
            raise ValueError("Word validation failed")
    elif spec.format == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(stream, data_only=False)
        try:
            if book.sheetnames != [sheet.name for sheet in spec.sheets]:
                raise ValueError("Spreadsheet validation failed")
            for sheet in spec.sheets:
                if sheet.chart != "none" and not book[sheet.name]._charts:
                    raise ValueError("Spreadsheet chart was not retained")
        finally:
            book.close()
    elif spec.format == "pptx":
        from pptx import Presentation

        if len(Presentation(stream).slides) != len(spec.slides):
            raise ValueError("Presentation validation failed")
    elif spec.format == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(stream)
        if not reader.pages:
            raise ValueError("PDF has no pages")
        if spec.kind == "form" and set(reader.get_fields() or {}) != {f.name for f in spec.fields}:
            raise ValueError("Fillable PDF fields were not retained")
    elif spec.format == "json":
        if json.loads(data) != spec.data:
            raise ValueError("JSON validation failed")
    elif spec.format == "yaml":
        import yaml

        if yaml.safe_load(data) != spec.data:
            raise ValueError("YAML validation failed")
    elif spec.format in {"xml", "svg"}:
        from defusedxml.ElementTree import fromstring

        fromstring(data)
    elif spec.format == "png":
        from PIL import Image

        Image.open(stream).verify()
    elif spec.format in {"csv", "tsv"}:
        rows = list(
            csv.reader(
                io.StringIO(data.decode("utf-8-sig")),
                delimiter="," if spec.format == "csv" else "\t",
            )
        )
        if len(rows) != len(spec.sheets[0].rows):
            raise ValueError("Delimited table validation failed")
    elif spec.format == "html":
        from html.parser import HTMLParser

        HTMLParser().feed(data.decode("utf-8"))
    elif spec.format == "rtf":
        if not data.startswith(b"{\\rtf1") or not data.endswith(b"}"):
            raise ValueError("RTF validation failed")
    else:
        data.decode("utf-8")


def generate(spec: Spec) -> Generated:
    check_spec(spec)
    text = text_content(spec)
    warning = ""
    fmt = spec.format
    if fmt in {"docx", "xlsx", "pptx", "pdf"}:
        data = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx, "pdf": _pdf}[fmt](spec)
    elif fmt in {"svg", "png"}:
        if fmt == "svg":
            from reportlab.graphics import renderSVG

            data = renderSVG.drawToString(_drawing(spec)).encode("utf-8")
        else:
            data = pdf_previews(_pdf(spec))[0]
    elif fmt in {"csv", "tsv"}:
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter="," if fmt == "csv" else "\t")
        for row in spec.sheets[0].rows:
            writer.writerow(
                [
                    "'" + c
                    if isinstance(c, str) and c.lstrip().startswith(("=", "+", "-", "@"))
                    else c
                    for c in row
                ]
            )
        data = output.getvalue().encode("utf-8-sig")
        warning = (
            "CSV/TSV contain values only; formula-like text is escaped for spreadsheet safety."
        )
    elif fmt == "json":
        data = (
            json.dumps(spec.data, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
        ).encode()
    elif fmt == "yaml":
        import yaml

        data = yaml.safe_dump(spec.data, allow_unicode=True, sort_keys=False).encode()
    elif fmt == "xml":
        data = _xml(spec.data)
    elif fmt == "html":
        data = _html(spec).encode()
    elif fmt == "rtf":
        data = _rtf(text)
    else:
        data = (_markdown(spec) if fmt == "md" else text).encode()
    validate_output(spec, data)
    if fmt == "xlsx" and any(
        isinstance(c, str) and c.startswith("=") for s in spec.sheets for row in s.rows for c in row
    ):
        warning = (
            "Formulas checked for supported syntax; "
            "results require Excel/LibreOffice recalculation."
        )
    preview = data if fmt == "pdf" else _pdf(spec)
    previews = pdf_previews(preview)
    warning = (warning + " " + REGISTRY[fmt].preview + ". Previews show up to 8 pages.").strip()
    if len(data) > 20 * 1024 * 1024:
        raise ValueError("Generated file exceeds 20 MB")
    return Generated(spec.model_dump(), data, previews, text, warning)
