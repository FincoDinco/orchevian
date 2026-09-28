"""Exercise a desktop bundle with temporary data and a spawned fake model."""

from __future__ import annotations

import argparse
import json
import os
import socket
import sys
import threading
import traceback
from pathlib import Path
from tempfile import TemporaryDirectory


def _web_fixture_worker(connection, query):
    """Exercise spawned retrieval without network or provider credentials."""
    from llm_engine.services.web_retrieval import SearchHit, retrieve

    class Provider:
        def search(self, query):
            return [SearchHit("Smoke reference", "https://example.com/reference")]

    try:
        result = retrieve(
            query, Provider(),
            lambda url: (url, "<p>Web smoke reference content.</p>", "text/html"),
            threading.Event(), lambda message: connection.send(("progress", message)),
        )
        result["provider"] = "Smoke fixture"
        connection.send(("done", result))
    finally:
        connection.close()


def _runtime_worker(connection):
    """Load the bundled native runtimes where the app runs models: a spawned worker."""
    import platform

    from llm_engine.backends.gguf import GGUFBackend
    from llm_engine.backends.mlx import MLXBackend

    try:
        found = []
        for backend in (GGUFBackend(), MLXBackend()):
            ok, reason = backend.is_available()
            found.append((str(backend.name), ok, reason))
        import llama_cpp

        llama_cpp.llama_backend_init()
        info = llama_cpp.llama_print_system_info().decode()
        if sys.platform == "darwin" and platform.machine() == "arm64":
            import mlx.core as mx
            from mlx_lm.sample_utils import make_sampler  # noqa: F401
            from transformers import AutoTokenizer  # noqa: F401

            if (mx.array([1, 2]) + 1).sum().item() != 5:
                raise RuntimeError("MLX computed a wrong result")
            info += f" | MLX {mx.default_device()}"
        connection.send(("done", (found, info)))
    except BaseException:
        connection.send(("error", traceback.format_exc()))
    finally:
        connection.close()


def _check_runtimes(checks: list[str]) -> None:
    import multiprocessing
    import platform

    receiver, sender = multiprocessing.get_context("spawn").Pipe(duplex=False)
    process = multiprocessing.get_context("spawn").Process(
        target=_runtime_worker, args=(sender,), daemon=True,
    )
    process.start()
    sender.close()
    try:
        if not receiver.poll(60):
            raise RuntimeError("Runtime check did not answer")
        kind, payload = receiver.recv()
    finally:
        process.join(10)
        if process.is_alive():
            process.kill()
    if kind != "done":
        raise RuntimeError(f"Bundled runtime failed to load:\n{payload}")
    found, info = payload
    expected = {"gguf"} | ({"mlx"} if sys.platform == "darwin"
                           and platform.machine() == "arm64" else set())
    missing = {name: reason for name, ok, reason in found if name in expected and not ok}
    if missing:
        raise RuntimeError(f"Bundled runtimes unavailable: {missing}")
    checks.append(f"bundled runtimes {sorted(expected)} ({info.strip()})")


def _check_real_models(registry, session, checks: list[str]) -> None:
    """Optional: ORCHEVIAN_SMOKE_MODEL_DIR runs the smallest real model per runtime."""
    from llm_engine.domain.models import ChatTurn, GenerationParams

    models, _ = registry.list_models()
    smallest = {}
    for model in models:
        name = str(model.ref.backend)
        if name in {"gguf", "mlx"} and model.available and (
            name not in smallest or model.size_bytes < smallest[name].size_bytes
        ):
            smallest[name] = model
    if not smallest:
        raise RuntimeError("ORCHEVIAN_SMOKE_MODEL_DIR has no GGUF or MLX models")
    for name, model in sorted(smallest.items()):
        cancel = threading.Event()
        session.load(model.ref, cancel=cancel)
        stream = session.generate(
            [ChatTurn("user", "Reply with one word: hello")],
            GenerationParams(temperature=0, max_tokens=16), cancel,
        )
        try:
            text = "".join(stream)
        finally:
            stream.close()
            session.force_unload()
        if not text.strip():
            raise RuntimeError(f"{model.ref.id} generated nothing")
        checks.append(f"real {name} generation ({model.ref.id}: {text.strip()[:40]!r})")


def _exercise(root: Path, checks: list[str]) -> None:
    import httpx
    import markdown
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from PySide6.QtCore import QSettings
    from PySide6.QtGui import QPixmap
    from PySide6.QtWidgets import QApplication

    from llm_engine.backends.fake import FakeBackend
    from llm_engine.backends.process import ProcessBackend
    from llm_engine.backends.registry import BackendRegistry
    from llm_engine.backends.smoke import VisionSmokeBackend
    from llm_engine.logging import setup_logging
    from llm_engine.services.catalog import CatalogService
    from llm_engine.services.chat import ChatService
    from llm_engine.services.documents import DocumentService, parse_isolated
    from llm_engine.services.openai_api import ApiServerService
    from llm_engine.services.session import ModelSession
    from llm_engine.store.library import LibraryService
    from llm_engine.store.sqlite import SqliteStore
    from llm_manager_app.main_window import MainWindow

    logger = setup_logging(log_path=root / "engine.log")
    logger.info("Desktop smoke check started")
    app = QApplication(["orchevian-smoke-test"])
    registry = BackendRegistry([ProcessBackend(FakeBackend(), FakeBackend, load_timeout=15)])
    session = ModelSession(registry)
    try:
        with SqliteStore(root / "data.db") as store:
            if store.schema_version() != 7:
                raise RuntimeError("Database migrations did not complete")
            library = LibraryService(store)
            template = library.save_template(name="Smoke test", user_prompt="Hello")
            if template.name != "Smoke test":
                raise RuntimeError("Database write failed")
            checks.append("database migrations and writes")

            api = ApiServerService(
                ChatService(library, session), CatalogService(registry, session),
            )
            # Reserve an unused port briefly; startup reports a bind failure if it is taken.
            with socket.socket() as sock:
                sock.bind(("127.0.0.1", 0))
                port = sock.getsockname()[1]
            try:
                api.start(port=port)
                with httpx.Client(base_url=f"http://127.0.0.1:{port}", trust_env=False,
                                  headers={"Authorization": f"Bearer {api.api_key}"},
                                  timeout=25) as client:
                    response = client.get("/v1/models")
                    response.raise_for_status()
                    model = response.json()["data"][0]["id"]
                    body = {"model": model, "messages": [{"role": "user", "content": "Hi"}]}
                    response = client.post("/v1/chat/completions", json=body)
                    response.raise_for_status()
                    if response.json()["choices"][0]["message"]["content"] != "Hello world":
                        raise RuntimeError("Spawned model response did not match")
                    response = client.post("/v1/chat/completions", json=body | {"stream": True})
                    response.raise_for_status()
                    if "data: [DONE]" not in response.text:
                        raise RuntimeError("Streaming response did not finish")
                if library.list_conversations():
                    raise RuntimeError("API smoke check unexpectedly saved a conversation")
                checks.extend(["spawned model load and generation", "local API JSON and streaming"])
            finally:
                api.stop()
                session.force_unload()

            # Build real Office files, then read them in disposable parser processes.
            word = Document()
            word.add_paragraph("Desktop document smoke check")
            word.save(root / "brief.docx")
            book = Workbook()
            book.active.append(["Budget", 450])
            book.save(root / "budget.xlsx")
            book.close()
            (root / "notes.txt").write_text("Desktop text smoke check", encoding="utf-8")
            (root / "data.csv").write_text("Name,Value\nBudget,450", encoding="utf-8")
            cid = library.create_conversation().summary.id
            documents = DocumentService(store)
            for name in ("brief.docx", "budget.xlsx", "notes.txt", "data.csv"):
                document = documents.import_file(cid, root / name, threading.Event())
                if not document.segments or documents.original(cid, document.id) != (
                    root / name
                ).read_bytes():
                    raise RuntimeError(f"Bundled document import failed: {name}")
            # A blank PDF must reach the PDF parser and report the OCR limitation.
            import io

            from llm_engine.domain.errors import EngineError

            pdf = PdfWriter()
            pdf.add_blank_page(width=200, height=200)
            buffer = io.BytesIO()
            pdf.write(buffer)
            try:
                parse_isolated(buffer.getvalue(), ".pdf", threading.Event())
            except EngineError as exc:
                if "OCR/vision" not in str(exc):
                    raise
            else:
                raise RuntimeError("Blank PDF incorrectly imported as readable text")
            library.delete_conversation(cid)
            checks.append("isolated document readers and original storage")

            from PIL import Image

            from llm_engine.domain.models import ChatTurn, GenerationParams

            Image.new("RGB", (100, 100), "white").save(root / "picture.png")
            image_doc = documents.import_file(-99, root / "picture.png", threading.Event())
            if len(image_doc.images) != 1:
                raise RuntimeError("Picture import did not retain a preview")
            _, _, pages = parse_isolated(buffer.getvalue(), ".pdf", threading.Event(),
                                          include_visuals=True)
            if len(pages) != 1 or pages[0].location != "Page 1":
                raise RuntimeError("PDF image rendering did not complete")
            visual_backend = ProcessBackend(FakeBackend(), VisionSmokeBackend, load_timeout=15)
            visual_session = ModelSession(BackendRegistry([visual_backend]))
            try:
                visual_session.load(FakeBackend().list_models()[0].ref)
                data = documents.visual_data(-99, image_doc, 0)
                response = "".join(visual_session.generate(
                    [ChatTurn("user", "Read picture", images=(data,))],
                    GenerationParams(), threading.Event(),
                ))
                if response != "Visual smoke passed":
                    raise RuntimeError("Image generation did not complete")
            finally:
                visual_session.force_unload()
                documents.discard_private(-99)
            checks.append("image decoding, PDF rendering and spawned vision input")

            project = library.create_project("Shared files smoke test")
            shared = documents.import_project_file(
                project.id, root / "notes.txt", threading.Event(),
            )
            cid = library.create_conversation(project_id=project.id).summary.id
            store.add_message(cid, "user", "Read the project notes")
            if "Desktop text smoke check" not in documents.context(cid, "notes"):
                raise RuntimeError("Shared project context was not retrieved")
            library.delete_project(project.id)
            if documents.sources(cid)[0]["document_id"] != shared.id:
                raise RuntimeError("Project deletion lost saved reply sources")
            library.delete_conversation(cid)
            checks.append("shared project files and retained sources")

            from llm_engine.services.web_retrieval import retrieve_isolated
            from llm_engine.services.web_search import WebSearchService

            web = WebSearchService(store, retriever=lambda query, cancel, progress: (
                retrieve_isolated(query, cancel, progress, worker=_web_fixture_worker)
            ))
            cid = library.create_conversation().summary.id
            store.add_message(cid, "user", "Web smoke question")
            if web.context(cid, "Web smoke question", False, threading.Event(), lambda _: None):
                raise RuntimeError("Search-off unexpectedly supplied web context")
            context = web.context(cid, "Web smoke question", True,
                                  threading.Event(), lambda _: None)
            if "Web smoke reference content" not in context or not web.history(cid)[0]["sources"]:
                raise RuntimeError("Spawned web reading did not retain evidence")
            library.delete_conversation(cid)
            checks.append("isolated web retrieval and retained sources (offline fixture)")

            from llm_engine.services.artifacts import ArtifactService, generate_isolated

            specs = [
                {"name": "proposal", "format": "docx", "kind": "document",
                 "title": "Proposal", "blocks": [{"text": "Bundled document creator"}]},
                {"name": "budget", "format": "xlsx", "kind": "spreadsheet",
                 "sheets": [{"name": "Budget",
                             "rows": [["Item", "Quantity", "Unit USD", "Total USD"],
                                      ["Work", 2, 225, "=B2*C2"],
                                      ["Total", None, None, "=SUM(D2:D2)"]],
                             "chart": "bar"}]},
                {"name": "briefing", "format": "pptx", "kind": "slides",
                 "slides": [{"title": "Briefing", "bullets": ["Budget: $450"],
                             "notes": "Bundled slide writer"}]},
                {"name": "intake", "format": "pdf", "kind": "form", "title": "Intake",
                 "fields": [{"name": "name", "label": "Your name"}]},
                {"name": "data", "format": "yaml", "kind": "data", "data": {"budget": 450}},
            ]
            created = generate_isolated(specs, threading.Event())
            if len(created) != 5 or not all(item.previews for item in created):
                raise RuntimeError("Bundled generators did not validate and render all files")
            from openpyxl import load_workbook

            generated_book = load_workbook(io.BytesIO(created[1].data))
            try:
                if generated_book["Budget"]["D2"].value != "=B2*C2":
                    raise RuntimeError("Bundled spreadsheet arithmetic did not survive export")
            finally:
                generated_book.close()
            cid = library.create_conversation().summary.id
            artifacts = ArtifactService(store)
            stored = artifacts.publish(cid, created, [], threading.Event())
            artifacts.export(cid, [item.id for item in stored], root / "deliverables.zip")
            library.delete_conversation(cid)
            checks.append("isolated artifact generators, previews and ZIP export")

            rendered = markdown.markdown(
                "```python\nprint('ok')\n```",
                extensions=["fenced_code", "nl2br", "sane_lists", "tables"],
            )
            if "<code" not in rendered:
                raise RuntimeError("Markdown extensions did not load")
            icon = QPixmap(str(Path(__file__).parent / "assets" / "chevron-down.svg"))
            if icon.isNull():
                raise RuntimeError("Bundled SVG asset could not be rendered")
            checks.append("Markdown extensions and SVG assets")

            if QPixmap(str(Path(__file__).parent / "assets" / "orchevian.png")).isNull():
                raise RuntimeError("Bundled app icon could not be loaded")
            from llm_manager_app.secret_store import _vault

            vault = _vault()
            # A real bundle must carry a vault backend; tests force the null one.
            frozen = getattr(sys, "frozen", False)
            if vault is None and frozen and sys.platform in {"darwin", "win32"}:
                raise RuntimeError("No system password vault backend was bundled")
            backend = type(vault.get_keyring()).__module__ if vault else "owner-only file"
            # Checked without writing; the workspace below must not touch the real vault.
            if vault is not None:
                from keyring.backends import null

                vault.set_keyring(null.Keyring())
            checks.append(f"app icon and secret storage ({backend})")

            _check_runtimes(checks)
            if os.environ.get("ORCHEVIAN_SMOKE_MODEL_DIR"):
                real = BackendRegistry()
                try:
                    _check_real_models(real, ModelSession(real), checks)
                finally:
                    real.close()

            settings = QSettings(str(root / "gui.ini"), QSettings.Format.IniFormat)
            window = MainWindow(registry=registry, library=library, settings=settings)
            try:
                window.show()
                app.processEvents()
                if window.grab().isNull():
                    raise RuntimeError("Workspace did not render")
            finally:
                if not window.close():
                    raise RuntimeError("Workspace did not close")
                app.processEvents()
            checks.append("Qt workspace startup and shutdown")
    finally:
        registry.close()
        for handler in list(logger.handlers):
            logger.removeHandler(handler)
            handler.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--smoke-test", required=True, type=Path, metavar="REPORT_JSON")
    args = parser.parse_args()
    checks: list[str] = []
    report = {"ok": False, "frozen": bool(getattr(sys, "frozen", False)), "checks": checks}
    overrides = ("QT_QPA_PLATFORM", "ORCHEVIAN_CONFIG", "ORCHEVIAN_DB")
    previous = {name: os.environ.get(name) for name in overrides}
    try:
        with TemporaryDirectory(prefix="orchevian-smoke-") as temporary:
            root = Path(temporary)
            os.environ["QT_QPA_PLATFORM"] = "offscreen"
            os.environ["ORCHEVIAN_CONFIG"] = str(root / "config.json")
            os.environ["ORCHEVIAN_DB"] = str(root / "data.db")
            (root / "config.json").write_text(
                json.dumps({
                    "model_dir": os.environ.get("ORCHEVIAN_SMOKE_MODEL_DIR", str(root / "models")),
                    "api_port": 8080,
                }),
                encoding="utf-8",
            )
            _exercise(root, checks)
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    finally:
        for name, value in previous.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value
    args.smoke_test.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return 0 if report["ok"] else 1
