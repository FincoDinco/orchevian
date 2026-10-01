from __future__ import annotations

import io
import json
import re
import threading
import zipfile

import pytest

from llm_engine.artifacts.generators import REGISTRY, generate, validate_output
from llm_engine.artifacts.specs import Spec, parse_call
from llm_engine.backends.fake import FakeBackend
from llm_engine.backends.registry import BackendRegistry
from llm_engine.domain.errors import EngineError
from llm_engine.services.artifacts import ArtifactService, generate_isolated
from llm_engine.services.chat import ChatService
from llm_engine.services.documents import DocumentService
from llm_engine.services.session import ModelSession
from llm_engine.store.library import LibraryService
from llm_engine.store.sqlite import SqliteStore


def specification(fmt="docx", kind=None):
    kind = kind or REGISTRY[fmt].kinds[0]
    fields = {
        "document": {
            "blocks": [
                {"type": "heading", "text": "Project proposal"},
                {"text": "A proposal grounded in the supplied brief [source-1]."},
                {"type": "list", "items": ["Research", "Delivery"]},
                {"type": "table", "rows": [["Item", "Cost"], ["Design", 450]]},
            ]
        },
        "spreadsheet": {
            "sheets": [
                {
                    "name": "Budget",
                    "rows": [["Item", "Cost"], ["Design", 450], ["Total", "=SUM(B2:B2)"]],
                    "number_formats": {"B2": "$#,##0.00"},
                    "chart": "bar" if fmt == "xlsx" else "none",
                }
            ]
        },
        "slides": {
            "slides": [
                {
                    "title": "Proposal",
                    "bullets": ["Scope", "Budget: $450"],
                    "notes": "Based on the project brief [source-1].",
                }
            ]
        },
        "form": {
            "fields": [
                {"name": "full_name", "label": "Full name"},
                {"name": "consent", "label": "Consent", "type": "checkbox"},
            ]
        },
        "data": {"data": {"budget": 450, "approved": True, "items": ["Design", "Review"]}},
        "chart": {"chart": {"labels": ["Design", "Review"], "values": [450, 100]}},
        "diagram": {"steps": ["Gather requirements", "Prepare proposal", "Review and approve"]},
    }
    return Spec(name="deliverable", format=fmt, kind=kind, title="Project brief", **fields[kind])


@pytest.mark.parametrize(
    "fmt,kind", [(fmt, kind) for fmt, cap in REGISTRY.items() for kind in cap.kinds]
)
def test_every_advertised_generator_has_valid_native_output_and_rendered_preview(fmt, kind):
    from PIL import Image

    spec = specification(fmt, kind)
    result = generate(spec)
    validate_output(spec, result.data)
    assert result.text and result.previews
    for preview in result.previews:
        with Image.open(io.BytesIO(preview)) as image:
            assert image.width > 100 and image.height > 100
    if fmt in {"docx", "pptx", "xlsx"}:
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            assert "[Content_Types].xml" in archive.namelist()


def test_editable_slides_notes_form_fields_and_spreadsheet_formulas():
    from openpyxl import load_workbook
    from pptx import Presentation
    from pypdf import PdfReader

    deck = Presentation(io.BytesIO(generate(specification("pptx")).data))
    assert any(
        "Budget: $450" in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    assert "source-1" in deck.slides[0].notes_slide.notes_text_frame.text
    pdf = PdfReader(io.BytesIO(generate(specification("pdf", "form")).data))
    assert pdf.get_fields()["full_name"]["/FT"] == "/Tx"
    assert pdf.get_fields()["consent"]["/FT"] == "/Btn"
    generated = generate(specification("xlsx"))
    book = load_workbook(io.BytesIO(generated.data))
    assert book["Budget"]["B3"].value == "=SUM(B2:B2)"
    assert book["Budget"]["B2"].number_format == "$#,##0.00"
    assert book["Budget"]._charts
    assert book.calculation.fullCalcOnLoad
    assert "recalculation" in generated.warning
    book.close()


def chart_ranges(spec):
    with zipfile.ZipFile(io.BytesIO(generate(spec).data)) as archive:
        chart = archive.read("xl/charts/chart1.xml").decode()
    # Series title, then category and value ranges; Excel hides an axis unless delete="0".
    return re.findall(r"<(?:\w+:)?f>([^<]*)<", chart)[1:], chart.count('<delete val="0"/>')


@pytest.mark.parametrize("label", ["Total", "TOTAL USD", "Grand total", "Subtotal", "Totals"])
def test_spreadsheet_chart_plots_items_without_trailing_totals_and_shows_axes(label):
    spec = specification("xlsx")
    spec.sheets[0].rows = [
        ["Item", "Cost"], ["Design", 450], ["Review", 100], [label, "=SUM(B2:B3)"],
    ]
    assert chart_ranges(spec) == (["'Budget'!$A$2:$A$3", "'Budget'!$B$2:$B$3"], 2)


@pytest.mark.parametrize("rows,ranges", [
    ([["Item", "Cost"], ["Design", 450], ["Totally new", 100]], ["$A$2:$A$3", "$B$2:$B$3"]),
    ([["Item", "Cost"], ["Total", 450]], ["$A$2", "$B$2"]),  # A lone data row is still charted.
])
def test_spreadsheet_chart_keeps_item_rows_that_resemble_totals(rows, ranges):
    spec = specification("xlsx")
    spec.sheets[0].rows = rows
    assert chart_ranges(spec)[0] == [f"'Budget'!{cells}" for cells in ranges]


def test_spreadsheet_columns_fit_content_and_print_one_page_wide():
    from openpyxl import load_workbook

    spec = specification("xlsx")
    item = "A considerably longer inventory item name"
    spec.sheets[0].rows = [
        ["Item", "Quantity", "Unit USD", "Total USD", "Notes"],
        [item, 4, 12.5, "=B2*C2", "x" * 500],
    ]
    book = load_workbook(io.BytesIO(generate(spec).data))
    try:
        ws = book["Budget"]
        widths = [ws.column_dimensions[column].width for column in "ABCDE"]
        assert widths[0] >= len(item) and widths[4] == 60
        assert max(widths[1:4]) < 22  # The fixed width that pushed columns onto a second page.
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 0)
    finally:
        book.close()


@pytest.mark.parametrize(
    "bad",
    [
        '{"tool":"run_python","files":[]}',
        '{"tool":"create_documents","files":[]}',
        "Here is your file: {}",
        '{"tool":"create_documents","files":[{"name":"../escape","format":"txt",'
        '"kind":"document","blocks":[{"text":"bad"}]}]}',
        '{"tool":"create_documents","files":[{"name":"file","format":"exe",'
        '"kind":"document","blocks":[{"text":"bad"}]}]}',
    ],
)
def test_malformed_or_unsafe_tool_calls_rejected(bad):
    with pytest.raises(ValueError):
        parse_call(bad)


def test_complete_file_array_missing_only_outer_brace_is_accepted():
    spec = specification("docx")
    # Include delimiters and escaped quotes in strings to rule out naive counting.
    spec.blocks[1].text = 'Literal ] } and "quoted" content stays exactly as supplied.'
    complete = json.dumps({"tool": "create_documents", "files": [spec.model_dump()]})
    for text in (complete, complete[:-1], "```json\n" + complete[:-1] + "\n```"):
        assert parse_call(text).files[0] == spec


@pytest.mark.parametrize("ending", ["", "}", "]", '"', '"}'])
def test_envelope_completion_never_completes_truncated_file_content(ending):
    partial = ('{"tool":"create_documents","files":[{"name":"report.docx",'
               '"format":"docx","kind":"document","blocks":[{"text":"Incomplete')
    with pytest.raises(ValueError):
        parse_call(partial + ending)


@pytest.mark.parametrize("change", [
    lambda text: text[:-2],  # Missing files array and outer object terminators.
    lambda text: text + "{}",  # More than one value.
    lambda text: text[:-1] + ",}",  # Trailing comma.
    lambda text: "Here is your file: " + text,
    lambda text: text.replace('"tool":', '"tool":"create_documents","tool":', 1),
    lambda text: text.replace('"kind":', '"kind":"document","kind":', 1)[:-1],
])
def test_envelope_completion_rejects_ambiguous_or_other_malformed_calls(change):
    complete = json.dumps({"tool": "create_documents", "files": [specification().model_dump()]})
    with pytest.raises(ValueError):
        parse_call(change(complete))


@pytest.mark.parametrize(
    "formula",
    [
        '=WEBSERVICE("https://example.com")',
        "=SUM(A1",
        "=[external.xlsx]A1",
        "=Missing!A1",
        "=1+",
        "=SUM()",
        "=IF(A1)",
    ],
)
def test_unsafe_or_invalid_formulas_fail(formula):
    spec = specification("xlsx")
    spec.sheets[0].rows[2][1] = formula
    with pytest.raises(ValueError):
        generate(spec)


@pytest.mark.parametrize("formula", [
    "=B2*2", "=B2+10", "=B2/2-1", "=B2^2", "=-B2", "=B2%",
    "=IF(B2>0,ROUND(B2*1.2,2),0)", "=ABS(-B2)", "=(B2+1)*2",
])
def test_spreadsheet_arithmetic_survives_native_readback(formula):
    from openpyxl import load_workbook

    spec = specification("xlsx")
    spec.sheets[0].rows[2][1] = formula
    book = load_workbook(io.BytesIO(generate(spec).data))
    try:
        assert book["Budget"]["B3"].value == formula
    finally:
        book.close()


@pytest.mark.parametrize("formula", [
    "=B3*2", "=$B$3+1", "=SUM(B2:B3)", "='Budget'!B3", "=Budget!$B$3",
])
def test_direct_circular_formulas_identify_the_output_cell(formula):
    spec = specification("xlsx")
    spec.sheets[0].rows[2][1] = formula
    with pytest.raises(ValueError, match="Budget!B3 references itself"):
        generate(spec)


@pytest.mark.parametrize("formula", ["=B2*", "=+", "=B2+*2", "=IF(B2>,1,0)", "=(B2+)2"])
def test_incomplete_arithmetic_is_still_rejected(formula):
    spec = specification("xlsx")
    spec.sheets[0].rows[2][1] = formula
    with pytest.raises(ValueError):
        generate(spec)


def test_html_has_no_model_markup_and_csv_escapes_formulas():
    spec = specification("html")
    spec.blocks[1].text = '<script src="https://example.com/x.js">alert(1)</script>'
    data = generate(spec).data.decode()
    assert "<script" not in data and "&lt;script" in data
    assert "default-src 'none'" in data
    result = generate(specification("csv"))
    assert "'=SUM" in result.data.decode("utf-8-sig")


@pytest.mark.parametrize("fmt", ["docx", "pdf", "html", "md"])
@pytest.mark.parametrize("heading", ["Workshop proposal", "# Workshop proposal"])
def test_document_renders_leading_title_once_without_losing_body(fmt, heading):
    from docx import Document
    from pypdf import PdfReader

    spec = Spec(name="report", format=fmt, kind="document", title="Workshop proposal", blocks=[
        {"type": "heading", "text": heading},
        {"text": "The body remains."},
        {"type": "heading", "text": "Budget"},
    ])
    result = generate(spec)
    if fmt == "docx":
        text = "\n".join(p.text for p in Document(io.BytesIO(result.data)).paragraphs)
    elif fmt == "pdf":
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(result.data)).pages)
    elif fmt == "html":
        text = result.data.decode().split("<body>", 1)[1]
    else:
        text = result.data.decode()
    assert text.count("Workshop proposal") == 1
    assert "The body remains." in text and "Budget" in text
    assert len(spec.blocks) == 3  # Rendering does not rewrite the retained model specification.


def test_versions_restart_templates_project_retrieval_zip_and_safe_export(tmp_path):
    path = tmp_path / "data.db"
    with SqliteStore(path) as store:
        library = LibraryService(store)
        project = library.create_project("Proposal")
        cid = library.create_conversation(project_id=project.id).summary.id
        store.add_message(cid, "user", "Create a proposal and slides")
        service = ArtifactService(store)
        cancel = threading.Event()
        generated = [generate(specification(fmt)) for fmt in ("docx", "pptx")]
        sources = [{"name": "brief.txt", "source": "source-1", "text": "Budget 450"}]
        first = service.publish(cid, generated, sources, cancel)
        second = service.publish(cid, generated[:1], sources, cancel)
        assert [a.version for a in service.list(cid)] == [1, 1, 2]
        key = service.save_template(cid, [a.id for a in first], "Proposal kit")
        assert (key, "Proposal kit") in service.templates()
        output = tmp_path / "kit.zip"
        service.export(cid, [a.id for a in first], output)
        with zipfile.ZipFile(output) as archive:
            assert set(archive.namelist()) == {"deliverable.docx", "deliverable.pptx"}
        original = output.read_bytes()
        with pytest.raises(FileExistsError):
            service.export(cid, [first[0].id], output)
        assert output.read_bytes() == original
        service.add_to_project(cid, first[1].id, project.id)
        other = library.create_conversation(project_id=project.id).summary.id
        store.add_message(other, "user", "What is the budget?")
        assert "Budget: $450" in DocumentService(store).context(other, "budget")
        assert service.get(cid, first[0].id)[1] == sources
    with SqliteStore(path) as store:
        service = ArtifactService(store)
        assert service.get(cid, first[0].id)[0].data == generated[0].data
        assert service.get(cid, second[0].id)[0].data == generated[0].data
        library = LibraryService(store)
        library.delete_project(project.id)
        assert len(service.list(cid)) == 3
        library.delete_conversation(cid)
        assert service.list(cid) == []


def test_private_generation_has_no_persistent_rows_and_cancelled_publish_is_atomic(tmp_path):
    with SqliteStore(tmp_path / "data.db") as store:
        service = ArtifactService(store)
        generated = [generate(specification("txt"))]
        cancel = threading.Event()
        service.publish(-1, generated, [], cancel)
        assert service.list(-1)
        with store.locked() as conn:
            assert conn.execute("SELECT count(*) FROM artifacts").fetchone()[0] == 0
        with pytest.raises(EngineError):
            service.save_template(-1, [service.list(-1)[0].id], "Private")
        service.discard_private(-1)
        assert not service.list(-1)
        with pytest.raises(EngineError, match="stopped"):
            service.publish(-1, generated, [], cancel)
        cancel.set()
        with pytest.raises(EngineError, match="stopped"):
            service.publish(-2, generated, [], cancel)
        assert not service.list(-2)


def test_spawned_generators_cancel_and_validate_native_files():
    specs = [specification(fmt).model_dump() for fmt in ("docx", "xlsx", "pptx")]
    progress = []
    result = generate_isolated(specs, threading.Event(), progress.append)
    assert len(result) == 3 and len(progress) == 3
    cancel = threading.Event()

    def stop_after_start(_message):
        cancel.set()

    with pytest.raises(EngineError, match="stopped"):
        generate_isolated(specs * 3, cancel, stop_after_start)


def test_failed_multifile_batch_returns_no_partial_result():
    good = specification("docx").model_dump()
    bad = specification("xlsx").model_dump()
    bad["sheets"][0]["rows"][2][1] = "=UNSUPPORTED(B2)"
    with pytest.raises(EngineError, match="Supported formula"):
        generate_isolated([good, bad], threading.Event())


class StructuredBackend(FakeBackend):
    def __init__(self, outputs):
        super().__init__()
        self.outputs = iter(outputs)
        self.prompts = []

    def stream_generate(self, handle, messages, params, cancel):
        self.prompts.append(messages)
        yield next(self.outputs)


def test_chat_tool_repair_source_grounding_revision_and_normal_chat_isolation(tmp_path):
    initial = specification("docx")
    initial.blocks[1].text = "A proposal grounded in brief.txt."
    revised = initial.model_copy(deep=True)
    revised.blocks[1].text = "Revised budget: $550. Source: brief.txt."

    def call(spec):
        return json.dumps({"tool": "create_documents", "files": [spec.model_dump()]})

    # The revision reproduces the live model's complete files array without its outer }.
    backend = StructuredBackend([
        "invalid output", call(initial), call(revised)[:-1], "Normal reply",
    ])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        session = ModelSession(BackendRegistry([backend]))
        errors, done, progress = [], [], []
        chat = ChatService(
            library,
            session,
            on_error=lambda *args: errors.append(args),
            on_done=lambda *args, **kwargs: done.append(kwargs),
        )
        chat.on_artifact_progress = lambda cid, message: progress.append(message)
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        source = tmp_path / "brief.txt"
        source.write_text("Proposal budget is $450")
        chat.documents.import_file(cid, source, threading.Event())
        chat.send(cid, "Create a proposal from the brief", artifact_request={})
        chat._worker_thread.join(30)
        assert not chat._worker_thread.is_alive()
        assert not errors and len(chat.artifacts.list(cid)) == 1
        assert "Proposal budget is $450" in "\n".join(t.content for t in backend.prompts[0])
        assert "Repairing" in " ".join(progress)
        artifact = chat.artifacts.list(cid)[0]
        assert chat.artifacts.get(cid, artifact.id)[1]
        chat.send(cid, "Change the budget to $550", artifact_request={"revision": artifact.id})
        chat._worker_thread.join(30)
        assert not errors and len(chat.artifacts.list(cid)) == 2
        assert chat.artifacts.list(cid)[1].version == 2
        assert len(backend.prompts) == 3  # Envelope completion needs no model retry.
        chat.send(cid, "Thanks")
        chat._worker_thread.join(10)
        assert library.get_conversation(cid).messages[-1].content == "Normal reply"
        assert "create_documents" not in "\n".join(t.content for t in backend.prompts[-1])
        assert len(chat.artifacts.list(cid)) == 2
        session.force_unload()


@pytest.mark.parametrize("extra_file", [False, True])
@pytest.mark.parametrize("repair_succeeds", [False, True])
def test_revision_preserves_target_and_original_on_wrong_file_response(
    tmp_path, extra_file, repair_succeeds,
):
    original = specification("docx")
    revised = original.model_copy(deep=True)
    revised.blocks[1].text = "Changed budget: $550."
    unrelated = revised.model_copy(update={"name": "unrelated"})
    invalid = [revised, unrelated] if extra_file else [unrelated]
    backend = StructuredBackend([
        json.dumps({"tool": "create_documents", "files": [s.model_dump() for s in files]})
        for files in (invalid, [revised] if repair_succeeds else invalid)
    ])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        session = ModelSession(BackendRegistry([backend]))
        errors = []
        chat = ChatService(library, session, on_error=lambda *args: errors.append(args))
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        store.add_message(cid, "user", "Create a proposal")
        generated = generate(original)
        first = chat.artifacts.publish(cid, [generated], [], threading.Event())[0]
        try:
            chat.send(cid, "Change the budget to $550", artifact_request={"revision": first.id})
            chat._worker_thread.join(20)
            assert not chat._worker_thread.is_alive()
            assert len(backend.prompts) == 2
            assert "exactly one file named deliverable.docx" in backend.prompts[1][-1].content
            artifacts = chat.artifacts.list(cid)
            assert chat.artifacts.get(cid, first.id)[0].data == generated.data
            assert [a.name for a in artifacts] == ["deliverable.docx"] * len(artifacts)
            assert [a.version for a in artifacts] == ([1, 2] if repair_succeeds else [1])
            assert bool(errors) is not repair_succeeds
        finally:
            session.force_unload()


def test_missing_source_attribution_gets_one_repair_before_publish(tmp_path):
    original = specification("docx")
    original.blocks[1].text = "A proposal with no source reference."
    fixed = Spec.model_validate({**original.model_dump(), "blocks": [
        *original.model_dump()["blocks"], {"text": "Source: brief.txt"},
    ]})
    backend = StructuredBackend([
        json.dumps({"tool": "create_documents", "files": [s.model_dump()]})
        for s in (original, fixed)
    ])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        session = ModelSession(BackendRegistry([backend]))
        errors = []
        chat = ChatService(library, session, on_error=lambda *args: errors.append(args))
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        brief = tmp_path / "brief.txt"
        brief.write_text("Proposal budget is $450")
        chat.documents.import_file(cid, brief, threading.Event())
        try:
            chat.send(cid, "Create a proposal from the brief", artifact_request={})
            chat._worker_thread.join(20)
            assert not chat._worker_thread.is_alive() and not errors
            assert len(chat.artifacts.list(cid)) == 1
            assert "missing source attribution" in backend.prompts[1][-1].content
            assert '"type": "paragraph", "text": "Sources: brief.txt"' in (
                backend.prompts[1][-1].content
            )
        finally:
            session.force_unload()


def test_model_without_structured_output_reports_error_without_partial_files(tmp_path):
    backend = StructuredBackend(["I cannot make files", "Here is a report"])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        session = ModelSession(BackendRegistry([backend]))
        errors = []
        chat = ChatService(library, session, on_error=lambda *args: errors.append(args))
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        chat.send(cid, "Create a report", artifact_request={})
        chat._worker_thread.join(20)
        assert errors[0][1].code == "artifact_failed"
        assert len(backend.prompts) == 2
        assert not chat.artifacts.list(cid)
        assert len(library.get_conversation(cid).messages) == 1
        session.force_unload()


def test_batch_attribution_errors_are_collected_for_one_repair(tmp_path):
    initial = [specification(fmt) for fmt in ("docx", "pdf", "pptx")]
    fixed = [spec.model_copy(deep=True) for spec in initial]
    for spec in fixed:
        if spec.kind == "document":
            spec.blocks[1].text = "Sources: brief.txt"
        else:
            spec.slides[0].notes = "Sources: brief.txt"
    backend = StructuredBackend([
        json.dumps({"tool": "create_documents", "files": [s.model_dump() for s in batch]})
        for batch in (initial, fixed)
    ])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        session = ModelSession(BackendRegistry([backend]))
        errors = []
        chat = ChatService(library, session, on_error=lambda *args: errors.append(args))
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        brief = tmp_path / "brief.txt"
        brief.write_text("Proposal budget is $450")
        chat.documents.import_file(cid, brief, threading.Event())
        try:
            chat.send(cid, "Create a proposal, PDF and slides", artifact_request={})
            chat._worker_thread.join(20)
            assert not chat._worker_thread.is_alive() and not errors
            assert len(backend.prompts) == 2 and len(chat.artifacts.list(cid)) == 3
            repair = backend.prompts[1][-1].content
            for name in ("deliverable.docx", "deliverable.pdf", "deliverable.pptx"):
                assert f"{name} is missing source attribution" in repair
        finally:
            session.force_unload()


def test_repair_lists_distinct_errors_after_many_malformed_rows(tmp_path):
    valid = specification("xlsx").model_dump()
    invalid = specification("xlsx").model_dump()
    invalid["sheets"][0]["rows"] = [{"items": ["Bad", 1]} for _ in range(30)]
    invalid["sheets"][0]["labels"] = ["Not a sheet field"]
    invalid["sheets"][0]["values"] = [1]
    backend = StructuredBackend([
        json.dumps({"tool": "create_documents", "files": [invalid]}),
        json.dumps({"tool": "create_documents", "files": [valid]}),
    ])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        session = ModelSession(BackendRegistry([backend]))
        errors = []
        chat = ChatService(library, session, on_error=lambda *args: errors.append(args))
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        try:
            chat.send(cid, "Create a budget", artifact_request={})
            chat._worker_thread.join(20)
            assert not chat._worker_thread.is_alive()
            assert not errors
            assert len(chat.artifacts.list(cid)) == 1
            repair = backend.prompts[1][-1].content
            assert repair.count("files.sheets.rows:") == 1
            assert "files.sheets.labels:" in repair
            assert "files.sheets.values:" in repair
            assert "errors.pydantic.dev" not in repair
            assert "input_value" not in repair
        finally:
            session.force_unload()


class SingleSystemBackend(StructuredBackend):
    """Rejects a second system message, as Qwen's chat template does."""

    def stream_generate(self, handle, messages, params, cancel):
        if any(turn.role == "system" for turn in messages[1:]):
            raise EngineError("backend_unavailable", "System message must be at the beginning.")
        yield from super().stream_generate(handle, messages, params, cancel)


def test_files_can_be_created_later_in_a_chat_with_guidance(tmp_path):
    call = json.dumps({"tool": "create_documents", "files": [specification("docx").model_dump()]})
    backend = SingleSystemBackend(["Here are five names.", call])
    with SqliteStore(tmp_path / "data.db") as store:
        library = LibraryService(store)
        errors = []
        chat = ChatService(library, ModelSession(BackendRegistry([backend])),
                           on_error=lambda *args: errors.append(args))
        cid = library.create_conversation(model=backend.list_models()[0].ref).summary.id
        chat.set_system_prompt(cid, "Keep answers short.")
        chat.send(cid, "Suggest five puppy names")
        chat._worker_thread.join(30)
        chat.send(cid, "Put those names in a Word document", artifact_request={})
        chat._worker_thread.join(30)
        assert not errors
        assert len(chat.artifacts.list(cid)) == 1
        system = backend.prompts[-1][0]
        assert system.role == "system"
        assert "create_documents" in system.content and "Keep answers short." in system.content
        assert "Suggest five puppy names" in " ".join(t.content for t in backend.prompts[-1])
