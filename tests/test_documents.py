from __future__ import annotations

import io
import json
import threading
import time

import pytest

from llm_engine.backends.fake import FakeBackend
from llm_engine.backends.registry import BackendRegistry
from llm_engine.domain.errors import EngineError
from llm_engine.services.chat import ChatService
from llm_engine.services.documents import CONTEXT_CHARS, DocumentService, extract, parse_isolated
from llm_engine.services.session import ModelSession
from llm_engine.store.library import LibraryService
from llm_engine.store.sqlite import SqliteStore


def pdf_bytes(*, encrypted=False, blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (The launch budget is 450 dollars.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_read_pdf_locations_and_unreadable_cases():
    segments, _ = extract(pdf_bytes(), ".pdf")
    assert segments[0][0] == "Page 1"
    assert "450 dollars" in segments[0][1]
    with pytest.raises(ValueError, match="Password-protected"):
        extract(pdf_bytes(encrypted=True), ".pdf")
    with pytest.raises(ValueError, match="OCR/vision"):
        extract(pdf_bytes(blank=True), ".pdf")


def test_word_headings_and_tables():
    from docx import Document

    doc = Document()
    doc.add_heading("Launch", 1)
    doc.add_paragraph("The budget is 450 dollars.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cost"
    table.cell(0, 1).text = "450"
    output = io.BytesIO()
    doc.save(output)
    segments, warning = extract(output.getvalue(), ".docx")
    assert any("Launch" in location and "450 dollars" in text for location, text in segments)
    assert any("table" in location and "Cost | 450" in text for location, text in segments)
    assert "images" in warning


def test_excel_cells_and_formula_limitations():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Cost", 450])
    sheet.append(["Total", "=B1"])
    output = io.BytesIO()
    book.save(output)
    segments, warning = extract(output.getvalue(), ".xlsx")
    assert any("Budget" in location and "B1: 450" in text for location, text in segments)
    assert any("B2: =B1" in text for _, text in segments)
    assert "not calculated" in warning


def test_text_csv_and_extraction_limits():
    segments, _ = extract(b'Name,Notes\nAlpha,"one,two"', ".csv")
    assert segments[-1] == ("Row 2", "Alpha | one,two")
    segments, warning = extract(("budget " * 40000).encode(), ".md")
    assert sum(len(text) for _, text in segments) <= 200000
    assert "limit reached" in warning
    with pytest.raises(ValueError, match="binary"):
        extract(b"\x00abc", ".txt")
    with pytest.raises(ValueError, match="UTF-8"):
        extract(b"\xff\xff", ".txt")


def test_isolated_parser_cancellation():
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(EngineError, match="cancelled"):
        parse_isolated(b"hello", ".txt", cancel)
    cancel.clear()
    assert parse_isolated(b"hello", ".txt", cancel)[0] == [("Line 1", "hello")]
    with pytest.raises(EngineError, match="too long"):
        parse_isolated(b"hello", ".txt", cancel, timeout=-1)


@pytest.fixture
def stack(tmp_path):
    store = SqliteStore(tmp_path / "data.db")
    library = LibraryService(store)
    service = DocumentService(store)
    yield store, library, service
    store.close()


def test_originals_survive_source_removal_and_database_reopen(stack, tmp_path):
    store, library, service = stack
    cid = library.create_conversation().summary.id
    path = tmp_path / "brief.txt"
    path.write_text("Launch budget 450")
    doc = service.import_file(cid, path, threading.Event())
    path.unlink()
    assert service.original(cid, doc.id) == b"Launch budget 450"
    assert service.context(cid, "budget") == ""  # Unsent drafts do not enter context.
    store.close()
    with SqliteStore(store.path) as reopened:
        assert DocumentService(reopened).list(cid) == [doc]
        LibraryService(reopened).delete_conversation(cid)
        with reopened.locked() as conn:
            assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0


def test_chat_commits_drafts_and_supplies_citable_scoped_excerpts(stack, tmp_path):
    store, library, _ = stack

    class RecordingBackend(FakeBackend):
        prompts = []

        def stream_generate(self, handle, messages, params, cancel):
            self.prompts.append(messages)
            yield "The budget is 450."

    fake = RecordingBackend()
    session = ModelSession(BackendRegistry([fake]))
    chat = ChatService(library, session)
    cid = library.create_conversation(model=fake.list_models()[0].ref).summary.id
    other = library.create_conversation(model=fake.list_models()[0].ref).summary.id
    path = tmp_path / "brief.txt"
    path.write_text("Launch budget 450")
    doc = chat.documents.import_file(cid, path, threading.Event())
    chat.send(cid, "What is the budget?")
    chat._worker_thread.join(5)
    assert not chat._worker_thread.is_alive()
    assert any("Launch budget 450" in turn.content for turn in fake.prompts[-1])
    assert chat.documents.list(cid)[0].sent
    sources = chat.documents.sources(cid)
    assert sources[0]["document_id"] == doc.id
    assert sources[0]["location"] == "Line 1"
    assert sum(len(json.dumps(item, ensure_ascii=False)) for item in sources) <= CONTEXT_CHARS
    chat.send(other, "What is the budget?")
    chat._worker_thread.join(5)
    assert not any("Launch budget 450" in turn.content for turn in fake.prompts[-1])
    chat.regenerate(cid)
    chat._worker_thread.join(5)
    assert any("Launch budget 450" in turn.content for turn in fake.prompts[-1])
    chat.send(cid, "Repeat the amount")
    chat._worker_thread.join(5)
    history = chat.documents.source_history(cid)
    assert [question for question, _ in history] == [
        "Repeat the amount", "What is the budget?"
    ]
    assert all(sources[0]["document_id"] == doc.id for _, sources in history)
    session.force_unload()


def test_failed_and_cancelled_imports_keep_existing_attachments(stack, tmp_path):
    _, library, service = stack
    cid = library.create_conversation().summary.id
    good = tmp_path / "brief.txt"
    good.write_text("Keep this brief")
    doc = service.import_file(cid, good, threading.Event())
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a Word archive")
    with pytest.raises(EngineError, match="Could not read"):
        service.import_file(cid, broken, threading.Event())
    cancelled = threading.Event()
    cancelled.set()
    with pytest.raises(EngineError, match="cancelled"):
        service.import_file(cid, good, cancelled)
    assert service.list(cid) == [doc]


def test_context_limit_and_draft_removal(stack, tmp_path):
    store, library, service = stack
    cid = library.create_conversation().summary.id
    path = tmp_path / "long.txt"
    path.write_text("\n".join(f"Budget line {i}: " + "details " * 150 for i in range(40)))
    doc = service.import_file(cid, path, threading.Event())
    message = store.add_message(cid, "user", "budget")
    with store.transaction() as conn:
        conn.execute("UPDATE documents SET message_id = ? WHERE id = ?", (message, doc.id))
    service.context(cid, "budget")
    sources = service.sources(cid)
    assert 0 < len(sources) < len(doc.segments)
    assert len(json.dumps(sources, ensure_ascii=False)) <= CONTEXT_CHARS
    service.remove_draft(cid, doc.id)
    assert service.list(cid)[0].sent
    draft = service.import_file(cid, path, threading.Event())
    service.remove_draft(cid, draft.id)
    assert len(service.list(cid)) == 1


def test_private_documents_never_persist_and_late_import_cannot_restore_them(stack, tmp_path):
    store, _, service = stack
    path = tmp_path / "private.txt"
    path.write_text("private document secret")
    doc = service.import_file(-1, path, threading.Event())
    service.commit_private(-1)
    assert "private document secret" in service.context(-1, "secret")
    assert service.original(-1, doc.id) == path.read_bytes()
    with store.locked() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM document_context").fetchone()[0] == 0
    service.discard_private(-1)
    assert service.list(-1) == [] and service.sources(-1) == []
    with pytest.raises(EngineError, match="cancelled"):
        service.import_file(-1, path, threading.Event())


def test_ui_upload_draft_isolation_and_remove(tmp_path):
    from test_app_shell import _qapp, _window

    app = _qapp()
    window, store, library = _window(tmp_path)
    first = library.create_conversation().summary.id
    second = library.create_conversation().summary.id
    panel = window._chat_view.attachments
    path = tmp_path / "brief.txt"
    path.write_text("Project brief")
    try:
        window._on_selected(first)
        panel.import_paths([str(path)])
        assert window._chat_view.composer()._importing
        window._on_selected(second)
        assert not window._chat_view.composer()._importing
        deadline = time.monotonic() + 5
        while panel.thread is not None and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.01)
        assert panel.thread is None
        assert panel.files.count() == 0
        window._on_selected(first)
        assert panel.files.count() == 1
        panel.remove_selected()
        assert panel.files.count() == 0
    finally:
        window.close()
        store.close()
