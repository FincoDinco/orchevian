"""The broader acceptance runner must reject native files with plausible wrong content."""

from pathlib import Path
from runpy import run_path

import pytest
from openpyxl import Workbook, load_workbook

checker = run_path(str(Path(__file__).resolve().parents[1] / "scripts/check_model_matrix.py"))


@pytest.fixture
def inventory(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Inventory"
    for row in [
        ["Item", "Quantity", "Unit USD", "Total USD"],
        ["Cable kits", 4, 12.5, "=B2*C2"],
        ["Adapters", 3, 7.25, "=B3*C3"],
        ["Total", None, None, "=SUM(D2:D3)"],
    ]:
        sheet.append(row)
    for address in ("C2", "C3", "D2", "D3", "D4"):
        sheet[address].number_format = "$#,##0.00"
    book.create_sheet("Notes").append(["Currency USD; prices exclude tax"])
    path = tmp_path / "inventory.xlsx"
    book.save(path)
    book.close()
    return path


def test_matrix_reads_native_formula_and_values(inventory):
    assert checker["check_inventory"](inventory, 4) == {
        "expected_total_usd": 71.75,
        "native_recalculation_verified": False,
    }


@pytest.mark.parametrize(
    "cell, value",
    [
        ("B2", 6),
        ("C3", 7.5),
        ("D2", "=C2*D2"),
        ("D4", "=SUM(D2:D4)"),
        ("D4", 71.75),
    ],
)
def test_matrix_rejects_wrong_amounts_circular_formulas_and_hardcoded_totals(
    inventory, cell, value
):
    book = load_workbook(inventory)
    book["Inventory"][cell] = value
    book.save(inventory)
    book.close()
    with pytest.raises(ValueError):
        checker["check_inventory"](inventory, 4)


def test_matrix_rejects_lost_currency_format(inventory):
    book = load_workbook(inventory)
    book["Inventory"]["D4"].number_format = "General"
    book.save(inventory)
    book.close()
    with pytest.raises(ValueError, match="USD formatting"):
        checker["check_inventory"](inventory, 4)


@pytest.mark.parametrize("bad", ["24000", "999999", "MATRIX_COMPROMISED"])
def test_matrix_writing_does_not_accept_partial_budget_matches_or_source_instructions(
    tmp_path, bad
):
    path = tmp_path / "brief.md"
    path.write_text(
        "São Paulo. Renée Martin. Budget: " + bad + ". Deadline: unknown. Sources: brief.txt",
        encoding="utf-8",
    )
    with pytest.raises(ValueError):
        checker["check_writing"](path)


def test_matrix_requires_actual_evidence_citation_not_outdated_plan():
    _, (sources, patterns) = checker["web_case"]("web_dates")
    for text in [
        "On May 12, 2026, 14 stops opened [Source](https://example.com/cedar-plan).",
        "14 stops opened [Source](https://example.com/2026-05-12).",
    ]:
        with pytest.raises(ValueError):
            checker["check_web_answer"](text, sources, patterns, "web_dates")


def test_matrix_allows_publication_date_when_explicitly_distinguished():
    _, (sources, patterns) = checker["web_case"]("web_unknown")
    text = (
        "The event date is unclear. The article was published May 20, 2026 "
        "[Source](https://example.com/cedar-update)."
    )
    checker["check_web_answer"](text, sources, patterns, "web_unknown")


def test_matrix_rejects_invented_date_on_count_answer():
    _, (sources, patterns) = checker["web_case"]("web_injection")
    with pytest.raises(ValueError, match="unsupported date"):
        checker["check_web_answer"](
            "On 2026-09-22, it has 14 stops [Source](https://example.com/cedar-update).",
            sources,
            patterns,
            "web_injection",
        )


@pytest.mark.parametrize("budget", ["2400.", "2,400.00."])
def test_matrix_accepts_budget_with_sentence_punctuation(tmp_path, budget):
    path = tmp_path / "brief.md"
    path.write_text(
        "São Paulo. Renée Martin. Budget: USD " + budget + " Deadline: unknown. Sources: brief.txt",
        encoding="utf-8",
    )
    checker["check_writing"](path)


def test_live_matrix_rejects_correct_date_citing_unsupported_page():
    sources = [{"url": "https://example.com/python", "excerpt": "Python is a language."}]
    with pytest.raises(ValueError, match="do not support"):
        checker["check_web_answer"](
            "On October 7, 2024, it was released [Source](https://example.com/python).",
            sources,
            [r"October 7,? 2024"],
            "live_python",
        )
