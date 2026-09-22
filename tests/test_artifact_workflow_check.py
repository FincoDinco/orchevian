"""Acceptance checks must reject readable files with wrong content or formatting."""

import json
from pathlib import Path
from runpy import run_path

import pytest

from llm_engine.artifacts.generators import generate
from llm_engine.artifacts.specs import Spec

checker = run_path(str(Path(__file__).resolve().parents[1] / "scripts/check_artifact_workflow.py"))
check_fixture_content = checker["check_fixture_content"]
BUDGET = checker["BUDGET"]
PROGRAM = checker["PROGRAM"]


@pytest.fixture
def fixture_files(tmp_path):
    def save(stage, name, kind, **content):
        folder = tmp_path / stage
        folder.mkdir(exist_ok=True)
        spec = Spec(name=name, format=name.rsplit(".", 1)[1], kind=kind,
                    title="Alder Community Workshop", **content)
        (folder / name).write_bytes(generate(spec).data)

    def blocks(count):
        return [
            {"text": f"October 24, 2026. Participants: {count}. Coordinator: Morgan Lee. "
                     "Purpose: teach residents basic bicycle maintenance. Source: project brief."},
            {"type": "list", "items": PROGRAM},
            {"type": "table", "rows": [["Item", "USD"], *[list(r) for r in BUDGET]]},
        ]

    save("primary", "proposal.docx", "document", blocks=blocks(40))
    save("revision", "proposal.docx", "document", blocks=blocks(50))
    save("primary", "summary.pdf", "document", blocks=blocks(40))
    save("primary", "budget.xlsx", "spreadsheet", sheets=[{
        "name": "Budget", "rows": [["Item", "Cost (USD)"],
                                    *[list(r) for r in BUDGET[:-1]], ["Total", "=SUM(B2:B4)"]],
        "number_formats": {f"B{n}": "$#,##0.00" for n in range(2, 6)}, "chart": "bar",
    }])
    save("primary", "briefing.pptx", "slides", slides=[
        {"title": "Overview", "bullets": ["October 24, 2026", "40 participants", "Morgan Lee",
                                          "Teach residents basic bicycle maintenance"],
         "notes": "Source: project brief."},
        {"title": "Budget", "bullets": [f"{name}: {value} USD" for name, value in BUDGET],
         "notes": "Source: project brief."},
        {"title": "Program", "bullets": PROGRAM, "notes": "Source: project brief."},
    ])
    save("extended", "registration.pdf", "form", fields=[
        {"name": "full_name", "label": "Full name"},
        {"name": "consent", "label": "Consent to attend", "type": "checkbox"},
    ])
    save("extended", "process.svg", "diagram", steps=PROGRAM)
    save("extended", "report.html", "document", blocks=blocks(40))
    save("extended", "budget.json", "data", data=dict(BUDGET))
    return tmp_path


def test_complete_native_fixture_passes(fixture_files):
    assert check_fixture_content(fixture_files) == []


def test_html_with_visible_markup_does_not_pass(fixture_files):
    path = fixture_files / "extended" / "report.html"
    path.write_text(path.read_text().replace("<table>", "&lt;table&gt;"), encoding="utf-8")
    assert "HTML displays literal markup instead of formatted content" in check_fixture_content(
        fixture_files,
    )


def test_diagram_checks_native_bytes_not_sidecar_spec(fixture_files):
    path = fixture_files / "extended" / "process.svg"
    path.with_suffix(".svg.spec.json").write_text(json.dumps({"steps": PROGRAM}))
    path.write_text(path.read_text().replace("Collect feedback", "Invented step"), encoding="utf-8")
    assert "Diagram did not preserve the program order" in check_fixture_content(fixture_files)


@pytest.mark.parametrize("data", [
    {"Room rental": 400, "Materials": 600, "Refreshments": 200, "Total": 1200},
    {name: str(value) for name, value in BUDGET},
    {"items": [{"name": name, "cost": value} for name, value in BUDGET[:-1]], "total": 999},
])
def test_wrong_or_string_json_amounts_fail(data):
    assert checker["json_budget_problems"](data)


def test_model_chosen_nested_json_shape_passes():
    data = {"budget": {"items": [{"label": name, "amount": value}
                                  for name, value in BUDGET[:-1]], "total": 1200}}
    assert checker["json_budget_problems"](data) == []


def test_wrong_slide_budget_fails_despite_valid_deck(fixture_files):
    from pptx import Presentation

    path = fixture_files / "primary" / "briefing.pptx"
    deck = Presentation(path)
    for shape in deck.slides[1].shapes:
        if shape.has_text_frame:
            shape.text = shape.text.replace("600", "999")
    deck.save(path)
    assert "Deck budget missing Room rental: 600" in check_fixture_content(fixture_files)


def test_wrong_html_participants_do_not_match_budget_amount(fixture_files):
    path = fixture_files / "extended" / "report.html"
    path.write_text(path.read_text().replace("Participants: 40", "Participants: 50"),
                    encoding="utf-8")
    assert "HTML missing 40" in check_fixture_content(fixture_files)
